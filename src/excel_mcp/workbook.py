import logging
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from .exceptions import WorkbookError

logger = logging.getLogger(__name__)

INVALID_WORKSHEET_NAME_CHARS = frozenset("\\/?*[]:")


def validate_worksheet_name(sheet_name: str) -> None:
    """Validate an Excel worksheet name with a stable, user-facing error."""
    if not isinstance(sheet_name, str) or not sheet_name.strip():
        raise WorkbookError("Worksheet name must not be empty")
    if len(sheet_name) > 31:
        raise WorkbookError("Worksheet name must be 31 characters or fewer")

    invalid = sorted(set(sheet_name) & INVALID_WORKSHEET_NAME_CHARS)
    if invalid:
        chars = " ".join(invalid)
        raise WorkbookError(
            f"Worksheet name contains invalid character(s): {chars}"
        )


def create_workbook(filepath: str, sheet_name: str = "Sheet1") -> dict[str, Any]:
    """Create a new Excel workbook with optional custom sheet name"""
    try:
        validate_worksheet_name(sheet_name)
        wb = Workbook()
        # Rename default sheet
        if "Sheet" in wb.sheetnames:
            sheet = wb["Sheet"]
            sheet.title = sheet_name
        else:
            wb.create_sheet(sheet_name)

        path = Path(filepath)
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(path))
        return {
            "message": f"Created workbook: {filepath}",
            "active_sheet": sheet_name,
            "workbook": wb
        }
    except Exception as e:
        logger.error(f"Failed to create workbook: {e}")
        raise WorkbookError(f"Failed to create workbook: {e!s}")

def get_or_create_workbook(filepath: str) -> Workbook:
    """Get existing workbook or create new one if it doesn't exist"""
    try:
        return load_workbook(filepath)
    except FileNotFoundError:
        return create_workbook(filepath)["workbook"]

def create_sheet(filepath: str, sheet_name: str) -> dict:
    """Create a new worksheet in the workbook if it doesn't exist."""
    wb = None
    try:
        validate_worksheet_name(sheet_name)
        wb = load_workbook(filepath)

        # Check if sheet already exists
        if sheet_name in wb.sheetnames:
            raise WorkbookError(f"Sheet {sheet_name} already exists")

        # Create new sheet
        wb.create_sheet(sheet_name)
        wb.save(filepath)
        return {"message": f"Sheet {sheet_name} created successfully"}
    except WorkbookError as e:
        logger.error(str(e))
        raise
    except Exception as e:
        logger.error(f"Failed to create sheet: {e}")
        raise WorkbookError(str(e))
    finally:
        if wb is not None:
            wb.close()

def get_workbook_info(filepath: str, include_ranges: bool = False) -> dict[str, Any]:
    """Get metadata about workbook including sheets, ranges, etc."""
    wb = None
    try:
        path = Path(filepath)
        if not path.exists():
            raise WorkbookError(f"File not found: {filepath}")

        wb = load_workbook(filepath, read_only=False)

        info = {
            "filename": path.name,
            "sheets": wb.sheetnames,
            "size": path.stat().st_size,
            "modified": path.stat().st_mtime
        }

        if include_ranges:
            # Add used ranges for each sheet
            ranges = {}
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                if ws.max_row > 0 and ws.max_column > 0:
                    # Skip completely empty sheets
                    if not (
                        ws.max_row == 1
                        and ws.max_column == 1
                        and ws.cell(1, 1).value is None
                    ):
                        ranges[sheet_name] = (
                            f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
                        )
            info["used_ranges"] = ranges

        return info

    except WorkbookError as e:
        logger.error(str(e))
        raise
    except Exception as e:
        logger.error(f"Failed to get workbook info: {e}")
        raise WorkbookError(str(e))
    finally:
        if wb is not None:
            wb.close()
