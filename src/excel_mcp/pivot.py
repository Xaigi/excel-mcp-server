from typing import Any
import uuid
import logging

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.styles import Font

from .cell_utils import parse_cell_range_strict
from .exceptions import ValidationError, PivotError

logger = logging.getLogger(__name__)

VALID_AGG_FUNCS = frozenset({"sum", "average", "count", "min", "max"})
EXCEL_SHEET_NAME_MAX = 31


def _parse_unqualified_range(data_range: str) -> tuple[str, str, int, int, int, int]:
    if not isinstance(data_range, str) or not data_range.strip():
        raise ValidationError("data_range must be a non-empty string")
    value = data_range.strip()
    if "!" in value:
        raise ValidationError(
            "data_range must not include a sheet qualifier; use sheet_name"
        )
    if ":" not in value:
        raise ValidationError("Data range must be in format 'A1:B2'")
    start_cell, end_cell = value.split(":", 1)
    start_row, start_col, end_row, end_col = parse_cell_range_strict(
        start_cell, end_cell
    )
    return start_cell, end_cell, start_row, start_col, end_row, end_col


def _clean_field_name(field: str) -> str:
    value = str(field).strip()
    for suffix in [" (sum)", " (average)", " (count)", " (min)", " (max)"]:
        if value.lower().endswith(suffix):
            return value[: -len(suffix)]
    return value


def _normalize_group_key(value: Any) -> str:
    """Normalize grouping keys so numeric 10 and string '10' match."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (int, float)):
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value)
    text = str(value).strip()
    if text == "":
        return ""
    try:
        as_float = float(text)
        if as_float.is_integer():
            return str(int(as_float))
        return str(as_float)
    except ValueError:
        return text


def _coerce_numeric(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        text = value.strip()
        if not text:
            return None
        try:
            return float(text)
        except ValueError:
            return None
    return None


def _is_non_empty(value: Any) -> bool:
    if value is None:
        return False
    if isinstance(value, str) and value.strip() == "":
        return False
    return True


def _require_unique_fields(fields: list[str], label: str) -> list[str]:
    if not isinstance(fields, list) or not fields:
        raise ValidationError(f"{label} must be a non-empty array of field names")
    cleaned = [_clean_field_name(field) for field in fields]
    if any(not name for name in cleaned):
        raise ValidationError(f"{label} field names must be non-empty")
    lowered = [name.lower() for name in cleaned]
    if len(set(lowered)) != len(lowered):
        raise ValidationError(f"{label} must not contain duplicate field names")
    return cleaned


def _read_rectangular_records(
    ws,
    start_row: int,
    start_col: int,
    end_row: int,
    end_col: int,
) -> list[dict[str, Any]]:
    if end_row <= start_row:
        raise PivotError("Source data must have a header row and at least one data row.")

    headers: list[str] = []
    for col in range(start_col, end_col + 1):
        value = ws.cell(row=start_row, column=col).value
        if value is None or (isinstance(value, str) and not value.strip()):
            raise PivotError("Source data headers must be non-empty")
        headers.append(str(value).strip())

    lowered = [h.lower() for h in headers]
    if len(set(lowered)) != len(lowered):
        raise PivotError("Source data headers must be unique")

    records: list[dict[str, Any]] = []
    for row in range(start_row + 1, end_row + 1):
        values = [
            ws.cell(row=row, column=col).value for col in range(start_col, end_col + 1)
        ]
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in values):
            continue
        records.append(dict(zip(headers, values)))

    if not records:
        raise PivotError("No data rows found after header.")
    return records


def _resolve_header(headers: list[str], field: str) -> str:
    wanted = field.lower()
    for header in headers:
        if header.lower() == wanted:
            return header
    raise ValidationError(
        f"Invalid field '{field}'. Available fields: {', '.join(headers)}"
    )


def create_pivot_table(
    filepath: str,
    sheet_name: str,
    data_range: str,
    rows: list[str],
    values: list[str],
    columns: list[str] | None = None,
    agg_func: str = "sum",
) -> dict[str, Any]:
    """Create a static summary worksheet from source data."""
    wb = None
    try:
        start_cell, end_cell, start_row, start_col, end_row, end_col = (
            _parse_unqualified_range(data_range)
        )

        agg = (agg_func or "").strip().lower()
        if agg not in VALID_AGG_FUNCS:
            raise ValidationError(
                f"Invalid aggregation function. Must be one of: "
                f"{', '.join(sorted(VALID_AGG_FUNCS))}"
            )

        cleaned_rows = _require_unique_fields(rows, "rows")
        cleaned_values = _require_unique_fields(values, "values")
        cleaned_columns = (
            _require_unique_fields(columns, "columns") if columns else []
        )

        pivot_sheet_name = f"{sheet_name}_pivot"
        if len(pivot_sheet_name) > EXCEL_SHEET_NAME_MAX:
            raise ValidationError(
                f"Pivot sheet name '{pivot_sheet_name}' exceeds Excel's "
                f"{EXCEL_SHEET_NAME_MAX}-character limit; use a shorter source sheet name"
            )

        wb = load_workbook(filepath)
        if sheet_name not in wb.sheetnames:
            raise ValidationError(f"Sheet '{sheet_name}' not found")

        ws = wb[sheet_name]
        data = _read_rectangular_records(
            ws, start_row, start_col, end_row, end_col
        )
        headers = list(data[0].keys())

        resolved_rows = [_resolve_header(headers, field) for field in cleaned_rows]
        resolved_values = [
            _resolve_header(headers, field) for field in cleaned_values
        ]
        resolved_columns = [
            _resolve_header(headers, field) for field in cleaned_columns
        ]

        # Unique combinations for row and column dimensions.
        row_field_values: dict[str, set[str]] = {field: set() for field in resolved_rows}
        col_field_values: dict[str, set[str]] = {
            field: set() for field in resolved_columns
        }
        for record in data:
            for field in resolved_rows:
                row_field_values[field].add(_normalize_group_key(record.get(field)))
            for field in resolved_columns:
                col_field_values[field].add(_normalize_group_key(record.get(field)))

        row_combinations = _get_combinations(row_field_values)
        col_combinations = (
            _get_combinations(col_field_values) if resolved_columns else [{}]
        )

        if pivot_sheet_name in wb.sheetnames:
            wb.remove(wb[pivot_sheet_name])
        pivot_ws = wb.create_sheet(pivot_sheet_name)

        # Header row
        current_col = 1
        for field in resolved_rows:
            cell = pivot_ws.cell(row=1, column=current_col, value=field)
            cell.font = Font(bold=True)
            current_col += 1

        value_headers: list[tuple[dict[str, str], str, str]] = []
        for col_combo in col_combinations:
            for value_field in resolved_values:
                if resolved_columns:
                    combo_label = " | ".join(
                        f"{field}={col_combo[field]}" for field in resolved_columns
                    )
                    header = f"{combo_label} | {value_field} ({agg})"
                else:
                    header = f"{value_field} ({agg})"
                cell = pivot_ws.cell(row=1, column=current_col, value=header)
                cell.font = Font(bold=True)
                value_headers.append((col_combo, value_field, header))
                current_col += 1

        total_cols = current_col - 1
        current_row = 2
        for row_combo in row_combinations:
            col = 1
            for field in resolved_rows:
                pivot_ws.cell(row=current_row, column=col, value=row_combo[field])
                col += 1

            for col_combo, value_field, _ in value_headers:
                filtered = _filter_data(data, row_combo, col_combo)
                try:
                    value = _aggregate_values(filtered, value_field, agg)
                except Exception as e:
                    raise PivotError(
                        f"Failed to aggregate values for field '{value_field}': {str(e)}"
                    ) from e
                pivot_ws.cell(row=current_row, column=col, value=value)
                col += 1
            current_row += 1

        total_rows = current_row - 1
        try:
            pivot_range = f"A1:{get_column_letter(total_cols)}{total_rows}"
            pivot_table = Table(
                displayName=f"PivotTable_{uuid.uuid4().hex[:8]}",
                ref=pivot_range,
            )
            style = TableStyleInfo(
                name="TableStyleMedium9",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=True,
            )
            pivot_table.tableStyleInfo = style
            pivot_ws.add_table(pivot_table)
        except Exception as e:
            raise PivotError(
                f"Failed to create pivot table formatting: {str(e)}"
            ) from e

        try:
            wb.save(filepath)
        except Exception as e:
            raise PivotError(f"Failed to save workbook: {str(e)}") from e

        data_range_str = (
            f"{get_column_letter(start_col)}{start_row}:"
            f"{get_column_letter(end_col)}{end_row}"
        )
        return {
            "message": "Summary table created successfully",
            "details": {
                "source_range": data_range_str,
                "pivot_sheet": pivot_sheet_name,
                "rows": resolved_rows,
                "columns": resolved_columns,
                "values": resolved_values,
                "aggregation": agg,
            },
        }
    except (ValidationError, PivotError):
        raise
    except Exception as e:
        logger.error(f"Failed to create pivot table: {e}")
        raise PivotError(str(e)) from e
    finally:
        if wb is not None:
            wb.close()


def _get_combinations(field_values: dict[str, set[str]]) -> list[dict[str, str]]:
    """Get all combinations of field values in deterministic order."""
    result: list[dict[str, str]] = [{}]
    for field, values in field_values.items():
        new_result: list[dict[str, str]] = []
        for combo in result:
            for value in sorted(values):
                new_combo = combo.copy()
                new_combo[field] = value
                new_result.append(new_combo)
        result = new_result
    return result


def _filter_data(
    data: list[dict[str, Any]],
    row_filters: dict[str, str],
    col_filters: dict[str, str],
) -> list[dict[str, Any]]:
    """Filter data based on normalized row and column filters."""
    result = []
    for record in data:
        matches = True
        for field, value in row_filters.items():
            if _normalize_group_key(record.get(field)) != value:
                matches = False
                break
        if not matches:
            continue
        for field, value in col_filters.items():
            if _normalize_group_key(record.get(field)) != value:
                matches = False
                break
        if matches:
            result.append(record)
    return result


def _aggregate_values(data: list[dict[str, Any]], field: str, agg_func: str) -> float:
    """Aggregate values using the specified function."""
    if agg_func == "count":
        return float(
            sum(1 for record in data if field in record and _is_non_empty(record[field]))
        )

    numbers: list[float] = []
    for record in data:
        if field not in record:
            continue
        coerced = _coerce_numeric(record[field])
        if coerced is not None:
            numbers.append(coerced)

    if not numbers:
        return 0.0

    if agg_func == "sum":
        return float(sum(numbers))
    if agg_func == "average":
        return float(sum(numbers) / len(numbers))
    if agg_func == "min":
        return float(min(numbers))
    if agg_func == "max":
        return float(max(numbers))
    raise ValidationError(f"Invalid aggregation function: {agg_func}")
