import logging
import re
from copy import copy
from typing import Any, Dict, Optional

from openpyxl import load_workbook
from openpyxl.styles import (
    PatternFill, Border, Side, Alignment, Protection, Font,
    Color
)
from openpyxl.formatting.rule import (
    ColorScaleRule, DataBarRule, IconSetRule,
    FormulaRule, CellIsRule
)

from .workbook import get_or_create_workbook
from .cell_utils import parse_cell_range, parse_cell_range_strict, validate_cell_reference
from .exceptions import ValidationError, FormattingError

logger = logging.getLogger(__name__)

ALLOWED_BORDER_STYLES = frozenset(
    {"thin", "medium", "thick", "double", "dashed", "dotted"}
)
ALLOWED_ALIGNMENTS = frozenset(
    {
        "general",
        "left",
        "center",
        "right",
        "fill",
        "justify",
        "centerContinuous",
        "distributed",
    }
)
_HEX_COLOR_RE = re.compile(r"^[0-9A-Fa-f]{6}([0-9A-Fa-f]{2})?$")


def normalize_hex_color(value: str, *, field_name: str) -> str:
    """Normalize a 6- or 8-character hex color without leading '#'."""
    if not isinstance(value, str) or not value:
        raise FormattingError(f"Invalid {field_name}: expected a hex color string")
    if value.startswith("#"):
        raise FormattingError(f"Invalid {field_name}: do not include '#'")
    if not _HEX_COLOR_RE.fullmatch(value):
        raise FormattingError(
            f"Invalid {field_name}: expected 6- or 8-character hexadecimal"
        )
    return value.upper() if len(value) == 8 else f"FF{value.upper()}"


def validate_format_options(format_options: Dict[str, Any]) -> Dict[str, Any]:
    """Validate and normalize Phase 3 format_options before download."""
    if not isinstance(format_options, dict) or not format_options:
        raise ValidationError("format_options must be a non-empty object")

    allowed = {
        "bold",
        "italic",
        "underline",
        "font_size",
        "font_color",
        "bg_color",
        "border_style",
        "border_color",
        "number_format",
        "alignment",
        "wrap_text",
    }
    unknown = sorted(set(format_options) - allowed)
    if unknown:
        raise ValidationError(
            f"Unknown format_options key(s): {', '.join(unknown)}"
        )

    cleaned: Dict[str, Any] = {}

    for key in ("bold", "italic", "underline", "wrap_text"):
        if key in format_options:
            value = format_options[key]
            if not isinstance(value, bool):
                raise ValidationError(f"{key} must be a boolean")
            cleaned[key] = value

    if "font_size" in format_options:
        value = format_options["font_size"]
        if isinstance(value, bool) or not isinstance(value, int) or value <= 0:
            raise ValidationError("font_size must be a positive integer")
        cleaned["font_size"] = value

    for key in ("font_color", "bg_color", "border_color"):
        if key in format_options:
            cleaned[key] = normalize_hex_color(
                format_options[key], field_name=key
            )

    if "border_style" in format_options:
        value = format_options["border_style"]
        if not isinstance(value, str) or value not in ALLOWED_BORDER_STYLES:
            raise ValidationError(
                "border_style must be one of: thin, medium, thick, double, dashed, dotted"
            )
        cleaned["border_style"] = value

    if "border_color" in cleaned and "border_style" not in cleaned:
        raise ValidationError("border_color requires border_style")

    if "number_format" in format_options:
        value = format_options["number_format"]
        if not isinstance(value, str) or not value:
            raise ValidationError("number_format must be a non-empty string")
        cleaned["number_format"] = value

    if "alignment" in format_options:
        value = format_options["alignment"]
        if not isinstance(value, str) or value not in ALLOWED_ALIGNMENTS:
            raise ValidationError(
                "alignment must be one of: general, left, center, right, fill, "
                "justify, centerContinuous, distributed"
            )
        cleaned["alignment"] = value

    return cleaned


def format_range(
    filepath: str,
    sheet_name: str,
    start_cell: str,
    end_cell: Optional[str] = None,
    bold: Optional[bool] = None,
    italic: Optional[bool] = None,
    underline: Optional[bool] = None,
    font_size: Optional[int] = None,
    font_color: Optional[str] = None,
    bg_color: Optional[str] = None,
    border_style: Optional[str] = None,
    border_color: Optional[str] = None,
    number_format: Optional[str] = None,
    alignment: Optional[str] = None,
    wrap_text: Optional[bool] = None,
    merge_cells: bool = False,
    protection: Optional[Dict[str, Any]] = None,
    conditional_format: Optional[Dict[str, Any]] = None,
    require_existing: bool = False,
    preserve_unset: bool = False,
) -> Dict[str, Any]:
    """Apply formatting to a range of cells."""
    wb = None
    try:
        if preserve_unset or require_existing:
            start_row, start_col, end_row, end_col = parse_cell_range_strict(
                start_cell, end_cell
            )
        else:
            if not validate_cell_reference(start_cell):
                raise ValidationError(f"Invalid start cell reference: {start_cell}")
            if end_cell and not validate_cell_reference(end_cell):
                raise ValidationError(f"Invalid end cell reference: {end_cell}")
            try:
                start_row, start_col, end_row, end_col = parse_cell_range(
                    start_cell, end_cell
                )
            except ValueError as e:
                raise ValidationError(f"Invalid cell range: {str(e)}") from e
            if end_row is None:
                end_row = start_row
            if end_col is None:
                end_col = start_col

        if require_existing:
            wb = load_workbook(filepath)
        else:
            wb = get_or_create_workbook(filepath)

        if sheet_name not in wb.sheetnames:
            raise ValidationError(f"Sheet '{sheet_name}' not found")

        sheet = wb[sheet_name]

        if preserve_unset:
            for row in range(start_row, end_row + 1):
                for col in range(start_col, end_col + 1):
                    cell = sheet.cell(row=row, column=col)

                    if any(
                        value is not None
                        for value in (bold, italic, underline, font_size, font_color)
                    ):
                        font_kwargs: Dict[str, Any] = {}
                        existing = cell.font
                        if existing is not None:
                            if existing.name is not None:
                                font_kwargs["name"] = existing.name
                            if existing.size is not None:
                                font_kwargs["size"] = existing.size
                            if existing.bold is not None:
                                font_kwargs["bold"] = existing.bold
                            if existing.italic is not None:
                                font_kwargs["italic"] = existing.italic
                            if existing.underline is not None:
                                font_kwargs["underline"] = existing.underline
                            if existing.color is not None:
                                font_kwargs["color"] = copy(existing.color)
                        if bold is not None:
                            font_kwargs["bold"] = bold
                        if italic is not None:
                            font_kwargs["italic"] = italic
                        if underline is not None:
                            font_kwargs["underline"] = (
                                "single" if underline else None
                            )
                        if font_size is not None:
                            font_kwargs["size"] = font_size
                        if font_color is not None:
                            color = (
                                font_color
                                if font_color.startswith("FF") and len(font_color) == 8
                                else (
                                    font_color
                                    if len(font_color) == 8
                                    else f"FF{font_color}"
                                )
                            )
                            font_kwargs["color"] = Color(rgb=color)
                        cell.font = Font(**font_kwargs)

                    if bg_color is not None:
                        color = (
                            bg_color
                            if len(bg_color) == 8
                            else f"FF{bg_color}"
                        )
                        cell.fill = PatternFill(
                            start_color=Color(rgb=color),
                            end_color=Color(rgb=color),
                            fill_type="solid",
                        )

                    if border_style is not None:
                        color_value = border_color or "000000"
                        color = (
                            color_value
                            if len(color_value) == 8
                            else f"FF{color_value}"
                        )
                        side = Side(style=border_style, color=Color(rgb=color))
                        cell.border = Border(
                            left=side, right=side, top=side, bottom=side
                        )

                    if alignment is not None or wrap_text is not None:
                        existing_align = cell.alignment
                        align_kwargs: Dict[str, Any] = {}
                        if existing_align is not None:
                            if existing_align.horizontal is not None:
                                align_kwargs["horizontal"] = existing_align.horizontal
                            if existing_align.vertical is not None:
                                align_kwargs["vertical"] = existing_align.vertical
                            if existing_align.wrap_text is not None:
                                align_kwargs["wrap_text"] = existing_align.wrap_text
                        if alignment is not None:
                            align_kwargs["horizontal"] = alignment
                        if wrap_text is not None:
                            align_kwargs["wrap_text"] = wrap_text
                        if "vertical" not in align_kwargs:
                            align_kwargs["vertical"] = "center"
                        cell.alignment = Alignment(**align_kwargs)

                    if number_format is not None:
                        cell.number_format = number_format
        else:
            # Legacy filepath-tool behavior: apply provided/default values.
            font_args = {
                "bold": bool(bold),
                "italic": bool(italic),
                "underline": "single" if underline else None,
            }
            if font_size is not None:
                font_args["size"] = font_size
            if font_color is not None:
                try:
                    font_color_norm = (
                        font_color
                        if font_color.startswith("FF")
                        else f"FF{font_color}"
                    )
                    font_args["color"] = Color(rgb=font_color_norm)
                except ValueError as e:
                    raise FormattingError(f"Invalid font color: {str(e)}") from e
            font = Font(**font_args)

            fill = None
            if bg_color is not None:
                try:
                    bg_color_norm = (
                        bg_color if bg_color.startswith("FF") else f"FF{bg_color}"
                    )
                    fill = PatternFill(
                        start_color=Color(rgb=bg_color_norm),
                        end_color=Color(rgb=bg_color_norm),
                        fill_type="solid",
                    )
                except ValueError as e:
                    raise FormattingError(
                        f"Invalid background color: {str(e)}"
                    ) from e

            border = None
            if border_style is not None:
                try:
                    border_color_value = border_color if border_color else "000000"
                    border_color_norm = (
                        border_color_value
                        if border_color_value.startswith("FF")
                        else f"FF{border_color_value}"
                    )
                    side = Side(
                        style=border_style,
                        color=Color(rgb=border_color_norm),
                    )
                    border = Border(
                        left=side, right=side, top=side, bottom=side
                    )
                except ValueError as e:
                    raise FormattingError(
                        f"Invalid border settings: {str(e)}"
                    ) from e

            align = None
            if alignment is not None or wrap_text:
                try:
                    align = Alignment(
                        horizontal=alignment,
                        vertical="center",
                        wrap_text=bool(wrap_text),
                    )
                except ValueError as e:
                    raise FormattingError(
                        f"Invalid alignment settings: {str(e)}"
                    ) from e

            protect = None
            if protection is not None:
                try:
                    protect = Protection(**protection)
                except ValueError as e:
                    raise FormattingError(
                        f"Invalid protection settings: {str(e)}"
                    ) from e

            for row in range(start_row, end_row + 1):
                for col in range(start_col, end_col + 1):
                    cell = sheet.cell(row=row, column=col)
                    cell.font = font
                    if fill is not None:
                        cell.fill = fill
                    if border is not None:
                        cell.border = border
                    if align is not None:
                        cell.alignment = align
                    if protect is not None:
                        cell.protection = protect
                    if number_format is not None:
                        cell.number_format = number_format

            if merge_cells and end_cell:
                try:
                    range_str = f"{start_cell}:{end_cell}"
                    sheet.merge_cells(range_str)
                except ValueError as e:
                    raise FormattingError(f"Failed to merge cells: {str(e)}") from e

            if conditional_format is not None:
                range_str = f"{start_cell}:{end_cell}" if end_cell else start_cell
                rule_type = conditional_format.get("type")
                if not rule_type:
                    raise FormattingError("Conditional format type not specified")

                params = conditional_format.get("params", {})

                if rule_type == "cell_is" and "fill" in params:
                    fill_params = params["fill"]
                    if isinstance(fill_params, dict):
                        try:
                            fill_color = fill_params.get("fgColor", "FFC7CE")
                            fill_color = (
                                fill_color
                                if fill_color.startswith("FF")
                                else f"FF{fill_color}"
                            )
                            params["fill"] = PatternFill(
                                start_color=fill_color,
                                end_color=fill_color,
                                fill_type="solid",
                            )
                        except ValueError as e:
                            raise FormattingError(
                                f"Invalid conditional format fill color: {str(e)}"
                            ) from e

                try:
                    if rule_type == "color_scale":
                        rule = ColorScaleRule(**params)
                    elif rule_type == "data_bar":
                        rule = DataBarRule(**params)
                    elif rule_type == "icon_set":
                        rule = IconSetRule(**params)
                    elif rule_type == "formula":
                        rule = FormulaRule(**params)
                    elif rule_type == "cell_is":
                        rule = CellIsRule(**params)
                    else:
                        raise FormattingError(
                            f"Invalid conditional format type: {rule_type}"
                        )

                    sheet.conditional_formatting.add(range_str, rule)
                except FormattingError:
                    raise
                except Exception as e:
                    raise FormattingError(
                        f"Failed to apply conditional formatting: {str(e)}"
                    ) from e

        wb.save(filepath)

        range_str = f"{start_cell}:{end_cell}" if end_cell else start_cell
        return {
            "message": f"Applied formatting to range {range_str}",
            "range": range_str,
        }

    except (ValidationError, FormattingError) as e:
        logger.error(str(e))
        raise
    except Exception as e:
        logger.error(f"Failed to apply formatting: {e}")
        raise FormattingError(str(e))
    finally:
        if wb is not None:
            wb.close()
