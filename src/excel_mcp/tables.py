import re
import uuid
import logging

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.worksheet.table import Table, TableStyleInfo

from .cell_utils import parse_cell_range_strict
from .exceptions import DataError, ValidationError

logger = logging.getLogger(__name__)

_TABLE_NAME_RE = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")
_CELL_LIKE_RE = re.compile(r"^[A-Za-z]+\d+$")

BUILTIN_TABLE_STYLES = frozenset(
    [f"TableStyleLight{i}" for i in range(1, 22)]
    + [f"TableStyleMedium{i}" for i in range(1, 29)]
    + [f"TableStyleDark{i}" for i in range(1, 12)]
)


def _parse_unqualified_range(data_range: str) -> tuple[int, int, int, int]:
    if not isinstance(data_range, str) or not data_range.strip():
        raise ValidationError("data_range must be a non-empty string")
    value = data_range.strip()
    if "!" in value:
        raise ValidationError(
            "data_range must not include a sheet qualifier; use sheet_name"
        )
    if ":" not in value:
        raise ValidationError("data_range must be in format 'A1:B2'")
    start_cell, end_cell = value.split(":", 1)
    return parse_cell_range_strict(start_cell, end_cell)


def _validate_table_name(table_name: str) -> str:
    if not isinstance(table_name, str) or not table_name.strip():
        raise DataError("table_name must be a non-empty string")
    name = table_name.strip()
    if " " in name:
        raise DataError("table_name must not contain spaces")
    if not _TABLE_NAME_RE.fullmatch(name):
        raise DataError(
            "table_name must start with a letter or underscore and contain "
            "only letters, digits, and underscores"
        )
    if _CELL_LIKE_RE.fullmatch(name):
        raise DataError("table_name must not look like a cell reference")
    return name


def _iter_workbook_tables(wb):
    for ws in wb.worksheets:
        for table in ws.tables.values():
            yield ws, table


def _ranges_overlap(left: CellRange, right: CellRange) -> bool:
    return not (
        left.max_row < right.min_row
        or left.min_row > right.max_row
        or left.max_col < right.min_col
        or left.min_col > right.max_col
    )


def create_excel_table(
    filepath: str,
    sheet_name: str,
    data_range: str,
    table_name: str | None = None,
    table_style: str = "TableStyleMedium9",
) -> dict:
    """Creates a native Excel table for the given data range."""
    wb = None
    try:
        start_row, start_col, end_row, end_col = _parse_unqualified_range(data_range)
        if end_row <= start_row:
            raise DataError(
                "data_range must include a header row and at least one data row"
            )
        if end_col < start_col:
            raise DataError("data_range end column cannot be before start column")

        if table_style not in BUILTIN_TABLE_STYLES:
            raise DataError(f"Unsupported table_style: {table_style}")

        wb = load_workbook(filepath)
        if sheet_name not in wb.sheetnames:
            raise DataError(f"Sheet '{sheet_name}' not found.")

        ws = wb[sheet_name]

        headers = []
        for col in range(start_col, end_col + 1):
            value = ws.cell(row=start_row, column=col).value
            if not isinstance(value, str) or not value.strip():
                raise DataError("Table headers must be non-empty strings")
            headers.append(value.strip())

        lowered = [h.lower() for h in headers]
        if len(set(lowered)) != len(lowered):
            raise DataError("Table headers must be unique")

        has_data_row = False
        for row in range(start_row + 1, end_row + 1):
            if any(
                ws.cell(row=row, column=col).value is not None
                for col in range(start_col, end_col + 1)
            ):
                has_data_row = True
                break
        if not has_data_row:
            raise DataError(
                "data_range must include a header row and at least one data row"
            )

        range_ref = (
            f"{get_column_letter(start_col)}{start_row}:"
            f"{get_column_letter(end_col)}{end_row}"
        )
        requested = CellRange(range_ref)
        for existing in ws.tables.values():
            if _ranges_overlap(CellRange(str(existing.ref)), requested):
                raise DataError(
                    f"data_range overlaps existing table '{existing.displayName}'"
                )

        if not table_name:
            table_name = f"Table_{uuid.uuid4().hex[:8]}"
        else:
            table_name = _validate_table_name(table_name)

        existing_names = {
            table.displayName.lower() for _, table in _iter_workbook_tables(wb)
        }
        if table_name.lower() in existing_names:
            raise DataError(f"Table name '{table_name}' already exists.")

        table = Table(displayName=table_name, ref=range_ref)
        table.tableStyleInfo = TableStyleInfo(
            name=table_style,
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(table)
        wb.save(filepath)

        return {
            "message": f"Successfully created table '{table_name}' in sheet '{sheet_name}'.",
            "table_name": table_name,
            "range": range_ref,
        }
    except (DataError, ValidationError):
        raise
    except Exception as e:
        logger.error(f"Failed to create table: {e}")
        raise DataError(str(e))
    finally:
        if wb is not None:
            wb.close()
