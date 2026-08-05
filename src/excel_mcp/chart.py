from typing import Any, Optional, Dict
import logging

from openpyxl import load_workbook
from openpyxl.chart import (
    BarChart,
    LineChart,
    PieChart,
    ScatterChart,
    AreaChart,
    Reference,
    Series,
)
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.legend import Legend
from openpyxl.chart.axis import ChartLines

from .cell_utils import parse_cell_range_strict, parse_cell_reference_strict
from .exceptions import ValidationError, ChartError

logger = logging.getLogger(__name__)

SUPPORTED_CHART_TYPES = frozenset({"line", "bar", "pie", "scatter", "area"})
_CHART_CLASSES = {
    "line": LineChart,
    "bar": BarChart,
    "pie": PieChart,
    "scatter": ScatterChart,
    "area": AreaChart,
}


def _parse_unqualified_range(data_range: str) -> tuple[int, int, int, int]:
    if not isinstance(data_range, str) or not data_range.strip():
        raise ValidationError("data_range must be a non-empty string")
    value = data_range.strip()
    if "!" in value:
        raise ValidationError(
            "data_range must not include a sheet qualifier; use sheet_name"
        )
    if ":" not in value:
        raise ValidationError("data_range must be in format 'A1:B2'")
    start_cell, end_cell = value.split(":", 1)
    return parse_cell_range_strict(start_cell, end_cell)


def create_chart_in_sheet(
    filepath: str,
    sheet_name: str,
    data_range: str,
    chart_type: str,
    target_cell: str,
    title: str = "",
    x_axis: str = "",
    y_axis: str = "",
    style: Optional[Dict] = None,
) -> dict[str, Any]:
    """Create chart in sheet with enhanced styling options."""
    if style is None:
        style = {"show_data_labels": True}
    else:
        style.setdefault("show_data_labels", True)

    wb = None
    try:
        start_row, start_col, end_row, end_col = _parse_unqualified_range(data_range)
        if end_row <= start_row:
            raise ValidationError(
                "data_range must include a header row and at least one data row"
            )
        if end_col <= start_col:
            raise ValidationError(
                "data_range must include at least two columns "
                "(categories plus one series)"
            )

        parse_cell_reference_strict(target_cell)

        chart_type_lower = (chart_type or "").strip().lower()
        ChartClass = _CHART_CLASSES.get(chart_type_lower)
        if not ChartClass:
            raise ValidationError(
                f"Unsupported chart type: {chart_type}. "
                f"Supported types: {', '.join(sorted(SUPPORTED_CHART_TYPES))}"
            )

        wb = load_workbook(filepath)
        if sheet_name not in wb.sheetnames:
            raise ValidationError(f"Sheet '{sheet_name}' not found")

        worksheet = wb[sheet_name]
        chart = ChartClass()
        chart.title = title or None
        if hasattr(chart, "x_axis") and x_axis:
            chart.x_axis.title = x_axis
        if hasattr(chart, "y_axis") and y_axis:
            chart.y_axis.title = y_axis

        try:
            if chart_type_lower == "scatter":
                for col in range(start_col + 1, end_col + 1):
                    x_values = Reference(
                        worksheet,
                        min_row=start_row + 1,
                        max_row=end_row,
                        min_col=start_col,
                    )
                    y_values = Reference(
                        worksheet,
                        min_row=start_row,
                        max_row=end_row,
                        min_col=col,
                    )
                    series = Series(y_values, x_values, title_from_data=True)
                    chart.series.append(series)
            else:
                data = Reference(
                    worksheet,
                    min_row=start_row,
                    max_row=end_row,
                    min_col=start_col + 1,
                    max_col=end_col,
                )
                cats = Reference(
                    worksheet,
                    min_row=start_row + 1,
                    max_row=end_row,
                    min_col=start_col,
                )
                chart.add_data(data, titles_from_data=True)
                chart.set_categories(cats)
        except (ValidationError, ChartError):
            raise
        except Exception as e:
            raise ChartError(f"Failed to create chart data references: {str(e)}") from e

        try:
            if style.get("show_legend", True):
                chart.legend = Legend()
                chart.legend.position = style.get("legend_position", "r")
            else:
                chart.legend = None

            if style.get("show_data_labels", False):
                data_labels = DataLabelList()
                dlo = (
                    style.get("data_label_options", {})
                    if isinstance(style.get("data_label_options", {}), dict)
                    else {}
                )

                def _opt(name: str, default: bool) -> bool:
                    return bool(dlo.get(name, default))

                data_labels.showVal = _opt("show_val", True)
                data_labels.showCatName = _opt("show_cat_name", False)
                data_labels.showSerName = _opt("show_ser_name", False)
                data_labels.showLegendKey = _opt("show_legend_key", False)
                data_labels.showPercent = _opt("show_percent", False)
                data_labels.showBubbleSize = _opt("show_bubble_size", False)
                chart.dataLabels = data_labels

            if style.get("grid_lines", False):
                if hasattr(chart, "x_axis"):
                    chart.x_axis.majorGridlines = ChartLines()
                if hasattr(chart, "y_axis"):
                    chart.y_axis.majorGridlines = ChartLines()
        except ChartError:
            raise
        except Exception as e:
            raise ChartError(f"Failed to apply chart style: {str(e)}") from e

        chart.width = 15
        chart.height = 7.5

        try:
            worksheet.add_chart(chart, target_cell)
        except Exception as e:
            raise ChartError(f"Failed to place chart: {str(e)}") from e

        try:
            wb.save(filepath)
        except Exception as e:
            raise ChartError(f"Failed to save workbook with chart: {str(e)}") from e

        return {
            "message": f"{chart_type_lower.capitalize()} chart created successfully",
            "details": {
                "type": chart_type_lower,
                "location": target_cell,
                "data_range": data_range,
            },
        }
    except (ValidationError, ChartError):
        raise
    except Exception as e:
        logger.error(f"Unexpected error creating chart: {e}")
        raise ChartError(f"Unexpected error creating chart: {str(e)}") from e
    finally:
        if wb is not None:
            wb.close()
