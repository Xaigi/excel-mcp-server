"""Remote workbook job orchestration using signed S3 URLs."""

from __future__ import annotations

import hashlib
import logging
import os
import shutil
import tempfile
import uuid
from typing import Any, Dict, List, Optional
from urllib.parse import urlparse

import httpx
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from excel_mcp.calculations import apply_formula
from excel_mcp.cell_utils import parse_cell_range, parse_cell_range_strict
from excel_mcp.data import write_data
from excel_mcp.exceptions import (
    CalculationError,
    DataError,
    FormattingError,
    SheetError,
    ValidationError,
    WorkbookError,
)
from excel_mcp.formatting import format_range, validate_format_options
from excel_mcp.sheet import (
    copy_range_operation,
    copy_sheet,
    delete_range_operation,
    delete_sheet,
    get_merged_ranges,
    merge_range,
    rename_sheet,
    unmerge_range,
)
from excel_mcp.validation import (
    validate_formula_in_cell_operation,
    validate_range_in_sheet_operation,
)
from excel_mcp.workbook import (
    create_sheet,
    create_workbook,
    get_workbook_info,
    validate_worksheet_name,
)

logger = logging.getLogger("excel-mcp")

DEFAULT_MAX_FILE_BYTES = 100 * 1024 * 1024
DEFAULT_TIMEOUT_SECONDS = 120.0
PREVIEW_MAX_ROWS = 10
ALLOWED_OPERATIONS = frozenset(
    {
        "create",
        "write",
        "read",
        "apply_formula",
        "validate_formula",
        "create_worksheet",
        "copy_worksheet",
        "rename_worksheet",
        "delete_worksheet",
        "format_range",
        "merge_cells",
        "unmerge_cells",
        "copy_range",
        "delete_range",
        "get_merged_cells",
        "validate_excel_range",
        "get_workbook_metadata",
    }
)
XLSX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)


class RemoteWorkbookError(Exception):
    """Safe, user-facing remote workbook failure."""

    def __init__(self, message: str):
        super().__init__(message)
        self.message = message


def _get_allowed_hosts() -> set[str]:
    raw = os.environ.get("EXCEL_MCP_ALLOWED_URL_HOSTS", "").strip()
    if not raw:
        return set()
    return {host.strip().lower() for host in raw.split(",") if host.strip()}


def _get_max_file_bytes() -> int:
    raw = os.environ.get(
        "EXCEL_MCP_MAX_FILE_BYTES", str(DEFAULT_MAX_FILE_BYTES)
    )
    try:
        value = int(raw)
    except ValueError as exc:
        raise RemoteWorkbookError(
            "Invalid EXCEL_MCP_MAX_FILE_BYTES configuration"
        ) from exc
    if value <= 0:
        raise RemoteWorkbookError(
            "EXCEL_MCP_MAX_FILE_BYTES must be positive"
        )
    return value


def _get_timeout_seconds() -> float:
    raw = os.environ.get(
        "EXCEL_MCP_REQUEST_TIMEOUT_SECONDS", str(DEFAULT_TIMEOUT_SECONDS)
    )
    try:
        value = float(raw)
    except ValueError as exc:
        raise RemoteWorkbookError(
            "Invalid EXCEL_MCP_REQUEST_TIMEOUT_SECONDS configuration"
        ) from exc
    if value <= 0:
        raise RemoteWorkbookError(
            "EXCEL_MCP_REQUEST_TIMEOUT_SECONDS must be positive"
        )
    return value


def validate_signed_url(url: str, *, purpose: str) -> None:
    """Validate a signed URL before any network request.

    Never include the URL value in raised errors.
    """
    if not url or not isinstance(url, str):
        raise RemoteWorkbookError(f"Missing {purpose} URL")

    try:
        parsed = urlparse(url)
    except Exception as exc:
        raise RemoteWorkbookError(f"Invalid {purpose} URL") from exc

    if parsed.scheme.lower() != "https":
        raise RemoteWorkbookError(f"{purpose} URL must use HTTPS")

    if parsed.username is not None or parsed.password is not None:
        raise RemoteWorkbookError(
            f"{purpose} URL must not contain credentials"
        )

    hostname = (parsed.hostname or "").lower()
    if not hostname:
        raise RemoteWorkbookError(f"{purpose} URL is missing a hostname")

    allowed = _get_allowed_hosts()
    if not allowed:
        raise RemoteWorkbookError(
            "EXCEL_MCP_ALLOWED_URL_HOSTS is not configured"
        )
    if hostname not in allowed:
        raise RemoteWorkbookError(
            f"{purpose} URL host is not allowed"
        )


def _create_http_client() -> httpx.Client:
    return httpx.Client(
        timeout=_get_timeout_seconds(),
        follow_redirects=False,
    )


def download_to_file(url: str, dest_path: str, max_bytes: int) -> int:
    """Stream-download a workbook to dest_path with a hard byte cap."""
    validate_signed_url(url, purpose="input download")

    with _create_http_client() as client:
        with client.stream("GET", url) as response:
            if response.is_redirect:
                raise RemoteWorkbookError(
                    "Redirects are not allowed for input download"
                )
            if response.status_code != 200:
                raise RemoteWorkbookError(
                    f"Input download failed with status {response.status_code}"
                )

            content_length = response.headers.get("Content-Length")
            if content_length is not None:
                try:
                    declared = int(content_length)
                except ValueError:
                    declared = -1
                if declared > max_bytes:
                    raise RemoteWorkbookError(
                        f"Input file exceeds maximum size of {max_bytes} bytes"
                    )

            total = 0
            with open(dest_path, "wb") as out:
                for chunk in response.iter_bytes():
                    if not chunk:
                        continue
                    total += len(chunk)
                    if total > max_bytes:
                        raise RemoteWorkbookError(
                            f"Input file exceeds maximum size of {max_bytes} bytes"
                        )
                    out.write(chunk)

    if total <= 0:
        raise RemoteWorkbookError("Downloaded workbook is empty")
    return total


def upload_file(
    url: str,
    file_path: str,
    upload_headers: Optional[Dict[str, str]] = None,
) -> int:
    """PUT a local workbook to a signed output URL."""
    validate_signed_url(url, purpose="output upload")

    size = os.path.getsize(file_path)
    max_bytes = _get_max_file_bytes()
    if size <= 0:
        raise RemoteWorkbookError("Output workbook is empty")
    if size > max_bytes:
        raise RemoteWorkbookError(
            f"Output file exceeds maximum size of {max_bytes} bytes"
        )

    headers = dict(upload_headers or {})
    with open(file_path, "rb") as body:
        with _create_http_client() as client:
            response = client.put(url, content=body, headers=headers)

    if response.is_redirect:
        raise RemoteWorkbookError(
            "Redirects are not allowed for output upload"
        )
    if response.status_code < 200 or response.status_code >= 300:
        raise RemoteWorkbookError(
            f"Output upload failed with status {response.status_code}"
        )
    return size


def _sha256_file(path: str) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _close_workbook_result(result: Any) -> None:
    workbook = None
    if isinstance(result, dict):
        workbook = result.get("workbook")
    if workbook is not None:
        try:
            workbook.close()
        except Exception:
            logger.debug("Failed to close workbook handle", exc_info=True)


def _build_upload_headers(
    upload_headers: Optional[Dict[str, str]],
    output_content_type: Optional[str],
) -> Dict[str, str]:
    headers = dict(upload_headers or {})
    if output_content_type and "Content-Type" not in headers:
        headers["Content-Type"] = output_content_type
    if "Content-Type" not in headers:
        headers["Content-Type"] = XLSX_CONTENT_TYPE
    return headers


def _upload_mutation_result(
    operation: str,
    output_upload_url: str,
    workbook_path: str,
    upload_headers: Optional[Dict[str, str]],
    output_content_type: Optional[str],
) -> Dict[str, Any]:
    headers = _build_upload_headers(upload_headers, output_content_type)
    size_bytes = upload_file(output_upload_url, workbook_path, headers)
    return {
        "success": True,
        "operation": operation,
        "output_uploaded": True,
        "size_bytes": size_bytes,
        "checksum_sha256": _sha256_file(workbook_path),
        "data": None,
    }


def read_rectangular_values(
    filepath: str,
    sheet_name: str,
    start_cell: str = "A1",
    end_cell: Optional[str] = None,
    preview_only: bool = False,
) -> Dict[str, Any]:
    """Read a rectangular range, preserving blank cells and blank rows."""
    wb = load_workbook(filepath, read_only=False, data_only=False)
    try:
        if sheet_name not in wb.sheetnames:
            raise DataError(f"Sheet '{sheet_name}' not found")

        ws = wb[sheet_name]

        if ":" in start_cell and end_cell is None:
            start_cell, end_cell = start_cell.split(":", 1)

        try:
            start_coords = parse_cell_range(f"{start_cell}:{start_cell}")
            if not start_coords or not all(
                coord is not None for coord in start_coords[:2]
            ):
                raise DataError(f"Invalid start cell reference: {start_cell}")
            start_row, start_col = start_coords[0], start_coords[1]
        except ValueError as exc:
            raise DataError(f"Invalid start cell format: {exc}") from exc

        if end_cell:
            try:
                end_coords = parse_cell_range(f"{end_cell}:{end_cell}")
                if not end_coords or not all(
                    coord is not None for coord in end_coords[:2]
                ):
                    raise DataError(f"Invalid end cell reference: {end_cell}")
                end_row, end_col = end_coords[0], end_coords[1]
            except ValueError as exc:
                raise DataError(f"Invalid end cell format: {exc}") from exc
        else:
            if (
                ws.max_row == 1
                and ws.max_column == 1
                and ws.cell(1, 1).value is None
            ):
                end_row, end_col = start_row, start_col
            else:
                end_row, end_col = ws.max_row, ws.max_column
                if start_cell.upper() == "A1":
                    start_row, start_col = ws.min_row, ws.min_column

        if start_row > end_row or start_col > end_col:
            raise DataError("Invalid range: start is after end")

        # When start is outside used area and no end_cell was provided, return empty.
        if end_cell is None and (
            start_row > ws.max_row or start_col > ws.max_column
        ):
            resolved_end = (
                f"{get_column_letter(start_col)}{start_row}"
            )
            return {
                "sheet_name": sheet_name,
                "start_cell": f"{get_column_letter(start_col)}{start_row}",
                "end_cell": resolved_end,
                "values": [],
            }

        if preview_only:
            end_row = min(end_row, start_row + PREVIEW_MAX_ROWS - 1)

        values: List[List[Any]] = []
        for row in range(start_row, end_row + 1):
            row_values: List[Any] = []
            for col in range(start_col, end_col + 1):
                row_values.append(ws.cell(row=row, column=col).value)
            values.append(row_values)

        return {
            "sheet_name": sheet_name,
            "start_cell": f"{get_column_letter(start_col)}{start_row}",
            "end_cell": f"{get_column_letter(end_col)}{end_row}",
            "values": values,
        }
    finally:
        wb.close()


def _require(condition: bool, message: str) -> None:
    if not condition:
        raise RemoteWorkbookError(message)


def execute_workbook_job(
    operation: str,
    request_id: Optional[str] = None,
    input_download_url: Optional[str] = None,
    output_upload_url: Optional[str] = None,
    output_content_type: Optional[str] = None,
    upload_headers: Optional[Dict[str, str]] = None,
    sheet_name: Optional[str] = None,
    start_cell: str = "A1",
    end_cell: Optional[str] = None,
    data: Optional[List[List[Any]]] = None,
    preview_only: bool = False,
    cell: Optional[str] = None,
    formula: Optional[str] = None,
    source_sheet: Optional[str] = None,
    target_sheet: Optional[str] = None,
    old_name: Optional[str] = None,
    new_name: Optional[str] = None,
    format_options: Optional[Dict[str, Any]] = None,
    source_start: Optional[str] = None,
    source_end: Optional[str] = None,
    target_start: Optional[str] = None,
    shift_direction: str = "up",
    include_ranges: bool = False,
) -> Dict[str, Any]:
    """Run one atomic workbook job using signed URLs."""
    req_id = request_id or str(uuid.uuid4())
    op = (operation or "").strip().lower()

    logger.info(
        "remote_workbook_job start request_id=%s operation=%s",
        req_id,
        op,
    )

    try:
        _require(op in ALLOWED_OPERATIONS, f"Unsupported operation: {operation}")
        max_bytes = _get_max_file_bytes()

        temp_root = tempfile.mkdtemp(prefix=f"excel-{req_id}-")
        try:
            workbook_path = os.path.join(temp_root, "workbook.xlsx")

            if op == "create":
                _require(
                    bool(output_upload_url),
                    "output_upload_url is required for create",
                )
                result = create_workbook(workbook_path)
                _close_workbook_result(result)
                response = _upload_mutation_result(
                    op,
                    output_upload_url,
                    workbook_path,
                    upload_headers,
                    output_content_type,
                )

            elif op == "write":
                _require(
                    bool(input_download_url),
                    "input_download_url is required for write",
                )
                _require(
                    bool(output_upload_url),
                    "output_upload_url is required for write",
                )
                _require(bool(sheet_name), "sheet_name is required for write")
                _require(
                    isinstance(data, list) and len(data) > 0,
                    "data is required for write and must be a non-empty array of arrays",
                )
                for row in data:
                    if not isinstance(row, list):
                        raise RemoteWorkbookError(
                            "data must be an array of arrays, not objects"
                        )

                download_to_file(input_download_url, workbook_path, max_bytes)
                write_data(workbook_path, sheet_name, data, start_cell or "A1")
                response = _upload_mutation_result(
                    op,
                    output_upload_url,
                    workbook_path,
                    upload_headers,
                    output_content_type,
                )

            elif op == "read":
                _require(
                    bool(input_download_url),
                    "input_download_url is required for read",
                )
                _require(bool(sheet_name), "sheet_name is required for read")

                download_to_file(input_download_url, workbook_path, max_bytes)
                read_data = read_rectangular_values(
                    workbook_path,
                    sheet_name,
                    start_cell or "A1",
                    end_cell,
                    preview_only=bool(preview_only),
                )
                response = {
                    "success": True,
                    "operation": "read",
                    "output_uploaded": False,
                    "data": read_data,
                }

            elif op == "apply_formula":
                _require(
                    bool(input_download_url),
                    "input_download_url is required for apply_formula",
                )
                _require(
                    bool(output_upload_url),
                    "output_upload_url is required for apply_formula",
                )
                _require(
                    bool(sheet_name),
                    "sheet_name is required for apply_formula",
                )
                _require(bool(cell), "cell is required for apply_formula")
                _require(
                    isinstance(formula, str) and bool(formula.strip()),
                    "formula is required for apply_formula",
                )
                validate_worksheet_name(sheet_name)

                download_to_file(input_download_url, workbook_path, max_bytes)
                apply_formula(
                    workbook_path,
                    sheet_name,
                    cell,
                    formula,
                    require_existing=True,
                )
                response = _upload_mutation_result(
                    op,
                    output_upload_url,
                    workbook_path,
                    upload_headers,
                    output_content_type,
                )

            elif op == "validate_formula":
                _require(
                    bool(input_download_url),
                    "input_download_url is required for validate_formula",
                )
                _require(
                    bool(sheet_name),
                    "sheet_name is required for validate_formula",
                )
                _require(bool(cell), "cell is required for validate_formula")
                _require(
                    isinstance(formula, str) and bool(formula.strip()),
                    "formula is required for validate_formula",
                )
                validate_worksheet_name(sheet_name)

                download_to_file(input_download_url, workbook_path, max_bytes)
                validation_data = validate_formula_in_cell_operation(
                    workbook_path,
                    sheet_name,
                    cell,
                    formula,
                )
                response = {
                    "success": True,
                    "operation": op,
                    "output_uploaded": False,
                    "data": validation_data,
                }

            elif op == "create_worksheet":
                _require(
                    bool(input_download_url),
                    "input_download_url is required for create_worksheet",
                )
                _require(
                    bool(output_upload_url),
                    "output_upload_url is required for create_worksheet",
                )
                _require(
                    bool(sheet_name),
                    "sheet_name is required for create_worksheet",
                )
                validate_worksheet_name(sheet_name)

                download_to_file(input_download_url, workbook_path, max_bytes)
                create_sheet(workbook_path, sheet_name)
                response = _upload_mutation_result(
                    op,
                    output_upload_url,
                    workbook_path,
                    upload_headers,
                    output_content_type,
                )

            elif op == "copy_worksheet":
                _require(
                    bool(input_download_url),
                    "input_download_url is required for copy_worksheet",
                )
                _require(
                    bool(output_upload_url),
                    "output_upload_url is required for copy_worksheet",
                )
                _require(
                    bool(source_sheet),
                    "source_sheet is required for copy_worksheet",
                )
                _require(
                    bool(target_sheet),
                    "target_sheet is required for copy_worksheet",
                )
                validate_worksheet_name(source_sheet)
                validate_worksheet_name(target_sheet)

                download_to_file(input_download_url, workbook_path, max_bytes)
                copy_sheet(workbook_path, source_sheet, target_sheet)
                response = _upload_mutation_result(
                    op,
                    output_upload_url,
                    workbook_path,
                    upload_headers,
                    output_content_type,
                )

            elif op == "rename_worksheet":
                _require(
                    bool(input_download_url),
                    "input_download_url is required for rename_worksheet",
                )
                _require(
                    bool(output_upload_url),
                    "output_upload_url is required for rename_worksheet",
                )
                _require(
                    bool(old_name),
                    "old_name is required for rename_worksheet",
                )
                _require(
                    bool(new_name),
                    "new_name is required for rename_worksheet",
                )
                validate_worksheet_name(old_name)
                validate_worksheet_name(new_name)

                download_to_file(input_download_url, workbook_path, max_bytes)
                rename_sheet(workbook_path, old_name, new_name)
                response = _upload_mutation_result(
                    op,
                    output_upload_url,
                    workbook_path,
                    upload_headers,
                    output_content_type,
                )

            elif op == "delete_worksheet":
                _require(
                    bool(input_download_url),
                    "input_download_url is required for delete_worksheet",
                )
                _require(
                    bool(output_upload_url),
                    "output_upload_url is required for delete_worksheet",
                )
                _require(
                    bool(sheet_name),
                    "sheet_name is required for delete_worksheet",
                )
                validate_worksheet_name(sheet_name)

                download_to_file(input_download_url, workbook_path, max_bytes)
                delete_sheet(workbook_path, sheet_name)
                response = _upload_mutation_result(
                    op,
                    output_upload_url,
                    workbook_path,
                    upload_headers,
                    output_content_type,
                )

            elif op == "format_range":
                _require(
                    bool(input_download_url),
                    "input_download_url is required for format_range",
                )
                _require(
                    bool(output_upload_url),
                    "output_upload_url is required for format_range",
                )
                _require(bool(sheet_name), "sheet_name is required for format_range")
                _require(bool(start_cell), "start_cell is required for format_range")
                cleaned_options = validate_format_options(format_options or {})
                validate_worksheet_name(sheet_name)
                parse_cell_range_strict(start_cell, end_cell)

                download_to_file(input_download_url, workbook_path, max_bytes)
                format_range(
                    filepath=workbook_path,
                    sheet_name=sheet_name,
                    start_cell=start_cell,
                    end_cell=end_cell,
                    require_existing=True,
                    preserve_unset=True,
                    **cleaned_options,
                )
                response = _upload_mutation_result(
                    op,
                    output_upload_url,
                    workbook_path,
                    upload_headers,
                    output_content_type,
                )

            elif op == "merge_cells":
                _require(
                    bool(input_download_url),
                    "input_download_url is required for merge_cells",
                )
                _require(
                    bool(output_upload_url),
                    "output_upload_url is required for merge_cells",
                )
                _require(bool(sheet_name), "sheet_name is required for merge_cells")
                _require(bool(start_cell), "start_cell is required for merge_cells")
                _require(bool(end_cell), "end_cell is required for merge_cells")
                validate_worksheet_name(sheet_name)
                parse_cell_range_strict(start_cell, end_cell)

                download_to_file(input_download_url, workbook_path, max_bytes)
                merge_range(workbook_path, sheet_name, start_cell, end_cell)
                response = _upload_mutation_result(
                    op,
                    output_upload_url,
                    workbook_path,
                    upload_headers,
                    output_content_type,
                )

            elif op == "unmerge_cells":
                _require(
                    bool(input_download_url),
                    "input_download_url is required for unmerge_cells",
                )
                _require(
                    bool(output_upload_url),
                    "output_upload_url is required for unmerge_cells",
                )
                _require(bool(sheet_name), "sheet_name is required for unmerge_cells")
                _require(bool(start_cell), "start_cell is required for unmerge_cells")
                _require(bool(end_cell), "end_cell is required for unmerge_cells")
                validate_worksheet_name(sheet_name)
                parse_cell_range_strict(start_cell, end_cell)

                download_to_file(input_download_url, workbook_path, max_bytes)
                unmerge_range(workbook_path, sheet_name, start_cell, end_cell)
                response = _upload_mutation_result(
                    op,
                    output_upload_url,
                    workbook_path,
                    upload_headers,
                    output_content_type,
                )

            elif op == "copy_range":
                _require(
                    bool(input_download_url),
                    "input_download_url is required for copy_range",
                )
                _require(
                    bool(output_upload_url),
                    "output_upload_url is required for copy_range",
                )
                _require(bool(sheet_name), "sheet_name is required for copy_range")
                _require(bool(source_start), "source_start is required for copy_range")
                _require(bool(source_end), "source_end is required for copy_range")
                _require(bool(target_start), "target_start is required for copy_range")
                validate_worksheet_name(sheet_name)
                if target_sheet:
                    validate_worksheet_name(target_sheet)

                download_to_file(input_download_url, workbook_path, max_bytes)
                copy_range_operation(
                    workbook_path,
                    sheet_name,
                    source_start,
                    source_end,
                    target_start,
                    target_sheet,
                )
                response = _upload_mutation_result(
                    op,
                    output_upload_url,
                    workbook_path,
                    upload_headers,
                    output_content_type,
                )

            elif op == "delete_range":
                _require(
                    bool(input_download_url),
                    "input_download_url is required for delete_range",
                )
                _require(
                    bool(output_upload_url),
                    "output_upload_url is required for delete_range",
                )
                _require(bool(sheet_name), "sheet_name is required for delete_range")
                _require(bool(start_cell), "start_cell is required for delete_range")
                _require(bool(end_cell), "end_cell is required for delete_range")
                validate_worksheet_name(sheet_name)
                parse_cell_range_strict(start_cell, end_cell)

                download_to_file(input_download_url, workbook_path, max_bytes)
                delete_range_operation(
                    workbook_path,
                    sheet_name,
                    start_cell,
                    end_cell,
                    shift_direction or "up",
                )
                response = _upload_mutation_result(
                    op,
                    output_upload_url,
                    workbook_path,
                    upload_headers,
                    output_content_type,
                )

            elif op == "get_merged_cells":
                _require(
                    bool(input_download_url),
                    "input_download_url is required for get_merged_cells",
                )
                _require(
                    bool(sheet_name),
                    "sheet_name is required for get_merged_cells",
                )
                validate_worksheet_name(sheet_name)

                download_to_file(input_download_url, workbook_path, max_bytes)
                merged = get_merged_ranges(workbook_path, sheet_name)
                response = {
                    "success": True,
                    "operation": op,
                    "output_uploaded": False,
                    "data": {
                        "sheet_name": sheet_name,
                        "merged_ranges": merged,
                    },
                }

            elif op == "validate_excel_range":
                _require(
                    bool(input_download_url),
                    "input_download_url is required for validate_excel_range",
                )
                _require(
                    bool(sheet_name),
                    "sheet_name is required for validate_excel_range",
                )
                _require(
                    bool(start_cell),
                    "start_cell is required for validate_excel_range",
                )
                validate_worksheet_name(sheet_name)
                parse_cell_range_strict(start_cell, end_cell)

                download_to_file(input_download_url, workbook_path, max_bytes)
                range_data = validate_range_in_sheet_operation(
                    workbook_path,
                    sheet_name,
                    start_cell,
                    end_cell,
                )
                response = {
                    "success": True,
                    "operation": op,
                    "output_uploaded": False,
                    "data": range_data,
                }

            else:  # get_workbook_metadata
                _require(
                    bool(input_download_url),
                    "input_download_url is required for get_workbook_metadata",
                )

                download_to_file(input_download_url, workbook_path, max_bytes)
                info = get_workbook_info(
                    workbook_path, include_ranges=bool(include_ranges)
                )
                metadata = {
                    "sheets": info.get("sheets", []),
                    "size_bytes": info.get("size", 0),
                }
                if include_ranges:
                    metadata["used_ranges"] = info.get("used_ranges", {})
                response = {
                    "success": True,
                    "operation": op,
                    "output_uploaded": False,
                    "data": metadata,
                }

            logger.info(
                "remote_workbook_job success request_id=%s operation=%s",
                req_id,
                op,
            )
            return response
        finally:
            shutil.rmtree(temp_root, ignore_errors=True)

    except RemoteWorkbookError as exc:
        logger.info(
            "remote_workbook_job failed request_id=%s operation=%s error=%s",
            req_id,
            op,
            exc.message,
        )
        return {"success": False, "error": exc.message}
    except (
        CalculationError,
        DataError,
        FormattingError,
        SheetError,
        ValidationError,
        WorkbookError,
    ) as exc:
        message = str(exc)
        logger.info(
            "remote_workbook_job failed request_id=%s operation=%s error=%s",
            req_id,
            op,
            message,
        )
        return {"success": False, "error": message}
    except httpx.HTTPError:
        logger.info(
            "remote_workbook_job failed request_id=%s operation=%s error=http_error",
            req_id,
            op,
        )
        return {
            "success": False,
            "error": "Network error while transferring workbook",
        }
    except Exception:
        logger.exception(
            "remote_workbook_job unexpected failure request_id=%s operation=%s",
            req_id,
            op,
        )
        return {
            "success": False,
            "error": "Unexpected error while executing workbook job",
        }
