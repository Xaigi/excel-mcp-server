import logging
from typing import Any, Dict, Optional
from copy import copy

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.styles import Font, Border, PatternFill, Side

from .cell_utils import (
    MAX_EXCEL_COL,
    MAX_EXCEL_ROW,
    parse_cell_range,
    parse_cell_range_strict,
    parse_cell_reference_strict,
)
from .exceptions import SheetError, ValidationError, WorkbookError
from .workbook import validate_worksheet_name

logger = logging.getLogger(__name__)


def _validate_sheet_name(sheet_name: str) -> None:
    try:
        validate_worksheet_name(sheet_name)
    except WorkbookError as exc:
        raise SheetError(str(exc)) from exc


def copy_sheet(filepath: str, source_sheet: str, target_sheet: str) -> Dict[str, Any]:
    """Copy a worksheet within the same workbook."""
    wb = None
    try:
        _validate_sheet_name(source_sheet)
        _validate_sheet_name(target_sheet)
        wb = load_workbook(filepath)
        if source_sheet not in wb.sheetnames:
            raise SheetError(f"Source sheet '{source_sheet}' not found")
            
        if target_sheet in wb.sheetnames:
            raise SheetError(f"Target sheet '{target_sheet}' already exists")
            
        source = wb[source_sheet]
        target = wb.copy_worksheet(source)
        target.title = target_sheet
        
        wb.save(filepath)
        return {"message": f"Sheet '{source_sheet}' copied to '{target_sheet}'"}
    except SheetError as e:
        logger.error(str(e))
        raise
    except Exception as e:
        logger.error(f"Failed to copy sheet: {e}")
        raise SheetError(str(e))
    finally:
        if wb is not None:
            wb.close()

def delete_sheet(filepath: str, sheet_name: str) -> Dict[str, Any]:
    """Delete a worksheet from the workbook."""
    wb = None
    try:
        _validate_sheet_name(sheet_name)
        wb = load_workbook(filepath)
        if sheet_name not in wb.sheetnames:
            raise SheetError(f"Sheet '{sheet_name}' not found")
            
        if len(wb.sheetnames) == 1:
            raise SheetError("Cannot delete the only sheet in workbook")
            
        del wb[sheet_name]
        wb.save(filepath)
        return {"message": f"Sheet '{sheet_name}' deleted"}
    except SheetError as e:
        logger.error(str(e))
        raise
    except Exception as e:
        logger.error(f"Failed to delete sheet: {e}")
        raise SheetError(str(e))
    finally:
        if wb is not None:
            wb.close()

def rename_sheet(filepath: str, old_name: str, new_name: str) -> Dict[str, Any]:
    """Rename a worksheet."""
    wb = None
    try:
        _validate_sheet_name(old_name)
        _validate_sheet_name(new_name)
        wb = load_workbook(filepath)
        if old_name not in wb.sheetnames:
            raise SheetError(f"Sheet '{old_name}' not found")
            
        if new_name in wb.sheetnames:
            raise SheetError(f"Sheet '{new_name}' already exists")
            
        sheet = wb[old_name]
        sheet.title = new_name
        wb.save(filepath)
        return {"message": f"Sheet renamed from '{old_name}' to '{new_name}'"}
    except SheetError as e:
        logger.error(str(e))
        raise
    except Exception as e:
        logger.error(f"Failed to rename sheet: {e}")
        raise SheetError(str(e))
    finally:
        if wb is not None:
            wb.close()

def format_range_string(start_row: int, start_col: int, end_row: int, end_col: int) -> str:
    """Format range string from row and column indices."""
    return f"{get_column_letter(start_col)}{start_row}:{get_column_letter(end_col)}{end_row}"

def copy_range(
    source_ws: Worksheet,
    target_ws: Worksheet,
    source_range: str,
    target_start: Optional[str] = None,
) -> None:
    """Copy range from source worksheet to target worksheet."""
    # Parse source range
    if ':' in source_range:
        source_start, source_end = source_range.split(':')
    else:
        source_start = source_range
        source_end = None
        
    src_start_row, src_start_col, src_end_row, src_end_col = parse_cell_range(
        source_start, source_end
    )

    if src_end_row is None:
        src_end_row = src_start_row
        src_end_col = src_start_col

    if target_start is None:
        target_start = source_start

    tgt_start_row, tgt_start_col, _, _ = parse_cell_range(target_start)

    for i, row in enumerate(range(src_start_row, src_end_row + 1)):
        for j, col in enumerate(range(src_start_col, src_end_col + 1)):
            source_cell = source_ws.cell(row=row, column=col)
            target_cell = target_ws.cell(row=tgt_start_row + i, column=tgt_start_col + j)

            target_cell.value = source_cell.value

            try:
                # Copy font
                font_kwargs = {}
                if hasattr(source_cell.font, 'name'):
                    font_kwargs['name'] = source_cell.font.name
                if hasattr(source_cell.font, 'size'):
                    font_kwargs['size'] = source_cell.font.size
                if hasattr(source_cell.font, 'bold'):
                    font_kwargs['bold'] = source_cell.font.bold
                if hasattr(source_cell.font, 'italic'):
                    font_kwargs['italic'] = source_cell.font.italic
                if hasattr(source_cell.font, 'color'):
                    font_color = None
                    if source_cell.font.color:
                        font_color = source_cell.font.color.rgb
                    font_kwargs['color'] = font_color
                target_cell.font = Font(**font_kwargs)

                # Copy border
                new_border = Border()
                for side in ['left', 'right', 'top', 'bottom']:
                    source_side = getattr(source_cell.border, side)
                    if source_side and source_side.style:
                        side_color = source_side.color.rgb if source_side.color else None
                        setattr(new_border, side, Side(
                            style=source_side.style,
                            color=side_color
                        ))
                target_cell.border = new_border

                # Copy fill
                if hasattr(source_cell, 'fill'):
                    fill_kwargs = {'patternType': source_cell.fill.patternType}
                    if hasattr(source_cell.fill, 'fgColor') and source_cell.fill.fgColor:
                        fg_color = None
                        if hasattr(source_cell.fill.fgColor, 'rgb'):
                            fg_color = source_cell.fill.fgColor.rgb
                        fill_kwargs['fgColor'] = fg_color
                    if hasattr(source_cell.fill, 'bgColor') and source_cell.fill.bgColor:
                        bg_color = None
                        if hasattr(source_cell.fill.bgColor, 'rgb'):
                            bg_color = source_cell.fill.bgColor.rgb
                        fill_kwargs['bgColor'] = bg_color
                    target_cell.fill = PatternFill(**fill_kwargs)

                # Copy number format and alignment
                if source_cell.number_format:
                    target_cell.number_format = source_cell.number_format
                if source_cell.alignment:
                    target_cell.alignment = source_cell.alignment

            except Exception:
                continue

def delete_range(worksheet: Worksheet, start_cell: str, end_cell: Optional[str] = None) -> None:
    """Delete contents and formatting of a range."""
    start_row, start_col, end_row, end_col = parse_cell_range(start_cell, end_cell)

    if end_row is None:
        end_row = start_row
        end_col = start_col

    for row in range(start_row, end_row + 1):
        for col in range(start_col, end_col + 1):
            cell = worksheet.cell(row=row, column=col)
            cell.value = None
            cell.font = Font()
            cell.border = Border()
            cell.fill = PatternFill()
            cell.number_format = "General"
            cell.alignment = None

def _ranges_overlap(left: CellRange, right: CellRange) -> bool:
    return not (
        left.max_row < right.min_row
        or left.min_row > right.max_row
        or left.max_col < right.min_col
        or left.min_col > right.max_col
    )


def _clear_cell(cell) -> None:
    cell.value = None
    cell.font = Font()
    cell.border = Border()
    cell.fill = PatternFill()
    cell.number_format = "General"
    cell.alignment = None


def merge_range(filepath: str, sheet_name: str, start_cell: str, end_cell: str) -> Dict[str, Any]:
    """Merge a range of cells."""
    wb = None
    try:
        _validate_sheet_name(sheet_name)
        wb = load_workbook(filepath)
        if sheet_name not in wb.sheetnames:
            raise SheetError(f"Sheet '{sheet_name}' not found")

        try:
            start_row, start_col, end_row, end_col = parse_cell_range_strict(
                start_cell, end_cell
            )
        except ValidationError as exc:
            raise SheetError(str(exc)) from exc

        range_string = format_range_string(start_row, start_col, end_row, end_col)
        worksheet = wb[sheet_name]
        requested = CellRange(range_string)
        for existing in worksheet.merged_cells.ranges:
            if str(existing).upper() == range_string.upper():
                raise SheetError(f"Range '{range_string}' is already merged")
            if _ranges_overlap(existing, requested):
                raise SheetError(
                    f"Range '{range_string}' overlaps existing merged range '{existing}'"
                )
        worksheet.merge_cells(range_string)
        wb.save(filepath)
        return {"message": f"Range '{range_string}' merged in sheet '{sheet_name}'"}
    except SheetError as e:
        logger.error(str(e))
        raise
    except Exception as e:
        logger.error(f"Failed to merge range: {e}")
        raise SheetError(str(e))
    finally:
        if wb is not None:
            wb.close()


def unmerge_range(filepath: str, sheet_name: str, start_cell: str, end_cell: str) -> Dict[str, Any]:
    """Unmerge a range of cells."""
    wb = None
    try:
        _validate_sheet_name(sheet_name)
        wb = load_workbook(filepath)
        if sheet_name not in wb.sheetnames:
            raise SheetError(f"Sheet '{sheet_name}' not found")

        worksheet = wb[sheet_name]

        try:
            start_row, start_col, end_row, end_col = parse_cell_range_strict(
                start_cell, end_cell
            )
        except ValidationError as exc:
            raise SheetError(str(exc)) from exc

        range_string = format_range_string(start_row, start_col, end_row, end_col)

        merged_ranges = worksheet.merged_cells.ranges
        target_range = range_string.upper()

        if not any(str(merged_range).upper() == target_range for merged_range in merged_ranges):
            raise SheetError(f"Range '{range_string}' is not merged")

        worksheet.unmerge_cells(range_string)
        wb.save(filepath)
        return {"message": f"Range '{range_string}' unmerged successfully"}
    except SheetError as e:
        logger.error(str(e))
        raise
    except Exception as e:
        logger.error(f"Failed to unmerge range: {e}")
        raise SheetError(str(e))
    finally:
        if wb is not None:
            wb.close()


def get_merged_ranges(filepath: str, sheet_name: str) -> list[str]:
    """Get merged cells in a worksheet."""
    wb = None
    try:
        _validate_sheet_name(sheet_name)
        wb = load_workbook(filepath)
        if sheet_name not in wb.sheetnames:
            raise SheetError(f"Sheet '{sheet_name}' not found")
        worksheet = wb[sheet_name]
        return [str(merged_range) for merged_range in worksheet.merged_cells.ranges]
    except SheetError as e:
        logger.error(str(e))
        raise
    except Exception as e:
        logger.error(f"Failed to get merged cells: {e}")
        raise SheetError(str(e))
    finally:
        if wb is not None:
            wb.close()


def copy_range_operation(
    filepath: str,
    sheet_name: str,
    source_start: str,
    source_end: str,
    target_start: str,
    target_sheet: Optional[str] = None
) -> Dict:
    """Copy a range of cells to another location."""
    wb = None
    try:
        _validate_sheet_name(sheet_name)
        if target_sheet is not None:
            _validate_sheet_name(target_sheet)

        wb = load_workbook(filepath)
        if sheet_name not in wb.sheetnames:
            raise ValidationError(f"Sheet '{sheet_name}' not found")

        dest_sheet = target_sheet or sheet_name
        if dest_sheet not in wb.sheetnames:
            raise ValidationError(f"Sheet '{dest_sheet}' not found")

        source_ws = wb[sheet_name]
        target_ws = wb[dest_sheet]

        try:
            start_row, start_col, end_row, end_col = parse_cell_range_strict(
                source_start, source_end
            )
            target_row, target_col = parse_cell_reference_strict(target_start)
        except ValidationError:
            raise

        row_span = end_row - start_row
        col_span = end_col - start_col
        target_end_row = target_row + row_span
        target_end_col = target_col + col_span
        if target_end_row > MAX_EXCEL_ROW or target_end_col > MAX_EXCEL_COL:
            raise ValidationError("Target range exceeds Excel worksheet bounds")

        snapshot = []
        for i in range(start_row, end_row + 1):
            for j in range(start_col, end_col + 1):
                source_cell = source_ws.cell(row=i, column=j)
                snapshot.append(
                    {
                        "value": source_cell.value,
                        "style": copy(source_cell._style) if source_cell.has_style else None,
                        "number_format": source_cell.number_format,
                        "alignment": copy(source_cell.alignment)
                        if source_cell.alignment is not None
                        else None,
                    }
                )

        idx = 0
        for i in range(start_row, end_row + 1):
            for j in range(start_col, end_col + 1):
                item = snapshot[idx]
                idx += 1
                target_cell = target_ws.cell(
                    row=target_row + (i - start_row),
                    column=target_col + (j - start_col),
                )
                target_cell.value = item["value"]
                if item["style"] is not None:
                    target_cell._style = item["style"]
                if item["number_format"] is not None:
                    target_cell.number_format = item["number_format"]
                if item["alignment"] is not None:
                    target_cell.alignment = item["alignment"]

        wb.save(filepath)
        return {"message": "Range copied successfully"}

    except (ValidationError, SheetError):
        raise
    except Exception as e:
        logger.error(f"Failed to copy range: {e}")
        raise SheetError(f"Failed to copy range: {str(e)}")
    finally:
        if wb is not None:
            wb.close()


def delete_range_operation(
    filepath: str,
    sheet_name: str,
    start_cell: str,
    end_cell: Optional[str] = None,
    shift_direction: str = "up"
) -> Dict[str, Any]:
    """Delete a range of cells and shift remaining cells within the range bounds."""
    wb = None
    try:
        _validate_sheet_name(sheet_name)
        wb = load_workbook(filepath)
        if sheet_name not in wb.sheetnames:
            raise SheetError(f"Sheet '{sheet_name}' not found")

        worksheet = wb[sheet_name]

        try:
            start_row, start_col, end_row, end_col = parse_cell_range_strict(
                start_cell, end_cell
            )
        except ValidationError as exc:
            raise SheetError(str(exc)) from exc

        if shift_direction not in ["up", "left"]:
            raise ValidationError(
                f"Invalid shift direction: {shift_direction}. Must be 'up' or 'left'"
            )

        range_string = format_range_string(start_row, start_col, end_row, end_col)
        max_row = worksheet.max_row
        max_col = worksheet.max_column

        if shift_direction == "up":
            height = end_row - start_row + 1
            for col in range(start_col, end_col + 1):
                for row in range(start_row, max_row + 1):
                    source_row = row + height
                    target = worksheet.cell(row=row, column=col)
                    if source_row <= max_row:
                        source = worksheet.cell(row=source_row, column=col)
                        target.value = source.value
                        if source.has_style:
                            target._style = copy(source._style)
                        target.number_format = source.number_format
                        target.alignment = (
                            copy(source.alignment)
                            if source.alignment is not None
                            else None
                        )
                    else:
                        _clear_cell(target)
        else:
            width = end_col - start_col + 1
            for row in range(start_row, end_row + 1):
                for col in range(start_col, max_col + 1):
                    source_col = col + width
                    target = worksheet.cell(row=row, column=col)
                    if source_col <= max_col:
                        source = worksheet.cell(row=row, column=source_col)
                        target.value = source.value
                        if source.has_style:
                            target._style = copy(source._style)
                        target.number_format = source.number_format
                        target.alignment = (
                            copy(source.alignment)
                            if source.alignment is not None
                            else None
                        )
                    else:
                        _clear_cell(target)

        wb.save(filepath)
        return {"message": f"Range {range_string} deleted successfully"}
    except (ValidationError, SheetError) as e:
        logger.error(str(e))
        raise
    except Exception as e:
        logger.error(f"Failed to delete range: {e}")
        raise SheetError(str(e))
    finally:
        if wb is not None:
            wb.close()

def insert_row(filepath: str, sheet_name: str, start_row: int, count: int = 1) -> Dict[str, Any]:
    """Insert one or more rows starting at the specified row."""
    try:
        wb = load_workbook(filepath)
        if sheet_name not in wb.sheetnames:
            raise SheetError(f"Sheet '{sheet_name}' not found")
            
        worksheet = wb[sheet_name]
        
        # Validate parameters
        if start_row < 1:
            raise ValidationError("Start row must be 1 or greater")
        if count < 1:
            raise ValidationError("Count must be 1 or greater")
            
        worksheet.insert_rows(start_row, count)
        wb.save(filepath)
        
        return {"message": f"Inserted {count} row(s) starting at row {start_row} in sheet '{sheet_name}'"}
    except (ValidationError, SheetError) as e:
        logger.error(str(e))
        raise
    except Exception as e:
        logger.error(f"Failed to insert rows: {e}")
        raise SheetError(str(e))

def insert_cols(filepath: str, sheet_name: str, start_col: int, count: int = 1) -> Dict[str, Any]:
    """Insert one or more columns starting at the specified column."""
    try:
        wb = load_workbook(filepath)
        if sheet_name not in wb.sheetnames:
            raise SheetError(f"Sheet '{sheet_name}' not found")
            
        worksheet = wb[sheet_name]
        
        # Validate parameters
        if start_col < 1:
            raise ValidationError("Start column must be 1 or greater")
        if count < 1:
            raise ValidationError("Count must be 1 or greater")
            
        worksheet.insert_cols(start_col, count)
        wb.save(filepath)
        
        return {"message": f"Inserted {count} column(s) starting at column {start_col} in sheet '{sheet_name}'"}
    except (ValidationError, SheetError) as e:
        logger.error(str(e))
        raise
    except Exception as e:
        logger.error(f"Failed to insert columns: {e}")
        raise SheetError(str(e))

def delete_rows(filepath: str, sheet_name: str, start_row: int, count: int = 1) -> Dict[str, Any]:
    """Delete one or more rows starting at the specified row."""
    try:
        wb = load_workbook(filepath)
        if sheet_name not in wb.sheetnames:
            raise SheetError(f"Sheet '{sheet_name}' not found")
            
        worksheet = wb[sheet_name]
        
        # Validate parameters
        if start_row < 1:
            raise ValidationError("Start row must be 1 or greater")
        if count < 1:
            raise ValidationError("Count must be 1 or greater")
        if start_row > worksheet.max_row:
            raise ValidationError(f"Start row {start_row} exceeds worksheet bounds (max row: {worksheet.max_row})")
            
        worksheet.delete_rows(start_row, count)
        wb.save(filepath)
        
        return {"message": f"Deleted {count} row(s) starting at row {start_row} in sheet '{sheet_name}'"}
    except (ValidationError, SheetError) as e:
        logger.error(str(e))
        raise
    except Exception as e:
        logger.error(f"Failed to delete rows: {e}")
        raise SheetError(str(e))

def delete_cols(filepath: str, sheet_name: str, start_col: int, count: int = 1) -> Dict[str, Any]:
    """Delete one or more columns starting at the specified column."""
    try:
        wb = load_workbook(filepath)
        if sheet_name not in wb.sheetnames:
            raise SheetError(f"Sheet '{sheet_name}' not found")
            
        worksheet = wb[sheet_name]
        
        # Validate parameters
        if start_col < 1:
            raise ValidationError("Start column must be 1 or greater")
        if count < 1:
            raise ValidationError("Count must be 1 or greater")
        if start_col > worksheet.max_column:
            raise ValidationError(f"Start column {start_col} exceeds worksheet bounds (max column: {worksheet.max_column})")
            
        worksheet.delete_cols(start_col, count)
        wb.save(filepath)
        
        return {"message": f"Deleted {count} column(s) starting at column {start_col} in sheet '{sheet_name}'"}
    except (ValidationError, SheetError) as e:
        logger.error(str(e))
        raise
    except Exception as e:
        logger.error(f"Failed to delete columns: {e}")
        raise SheetError(str(e))
