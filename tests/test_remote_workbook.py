import io
import os
import sys
import tempfile
import unittest
from unittest.mock import MagicMock, patch

_REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
_SRC = os.path.join(_REPO_ROOT, "src")
if _SRC not in sys.path:
    sys.path.insert(0, _SRC)

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

import excel_mcp.remote_workbook as remote
from excel_mcp.workbook import create_workbook


ALLOWED_HOST = "mimasa-workflows-local-dev.s3.ap-south-1.amazonaws.com"


def _url(path: str = "/file.xlsx") -> str:
    return f"https://{ALLOWED_HOST}{path}"


class _FakeStreamResponse:
    def __init__(
        self,
        content: bytes,
        status_code: int = 200,
        headers: dict | None = None,
        is_redirect: bool = False,
    ):
        self._content = content
        self.status_code = status_code
        self.headers = headers or {}
        self.is_redirect = is_redirect

    def iter_bytes(self):
        # Yield in small chunks so size-cap streaming can be tested.
        chunk_size = 8
        for i in range(0, len(self._content), chunk_size):
            yield self._content[i : i + chunk_size]

    def __enter__(self):
        return self

    def __exit__(self, exc_type, exc, tb):
        return False


class _FakeClient:
    def __init__(self, get_response=None, put_response=None):
        self.get_response = get_response
        self.put_response = put_response
        self.put_calls = []
        self.get_calls = []

    def stream(self, method, url):
        self.get_calls.append((method, url))
        return self.get_response

    def put(self, url, content=None, headers=None):
        # Consume the file-like content so callers behave normally.
        if hasattr(content, "read"):
            body = content.read()
        else:
            body = content
        self.put_calls.append(
            {"url": url, "headers": dict(headers or {}), "body": body}
        )
        return self.put_response

    def __enter__(self):
        return self

    def __exit__(self, exc_type, exc, tb):
        return False


class RemoteWorkbookTests(unittest.TestCase):
    def setUp(self):
        self._old_env = {
            "EXCEL_MCP_ALLOWED_URL_HOSTS": os.environ.get(
                "EXCEL_MCP_ALLOWED_URL_HOSTS"
            ),
            "EXCEL_MCP_MAX_FILE_BYTES": os.environ.get(
                "EXCEL_MCP_MAX_FILE_BYTES"
            ),
            "EXCEL_MCP_REQUEST_TIMEOUT_SECONDS": os.environ.get(
                "EXCEL_MCP_REQUEST_TIMEOUT_SECONDS"
            ),
        }
        os.environ["EXCEL_MCP_ALLOWED_URL_HOSTS"] = ALLOWED_HOST
        os.environ["EXCEL_MCP_MAX_FILE_BYTES"] = str(1024 * 1024)
        os.environ["EXCEL_MCP_REQUEST_TIMEOUT_SECONDS"] = "30"

    def tearDown(self):
        for key, value in self._old_env.items():
            if value is None:
                os.environ.pop(key, None)
            else:
                os.environ[key] = value

    def _xlsx_bytes(self) -> bytes:
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "sample.xlsx")
            create_workbook(path)
            with open(path, "rb") as f:
                return f.read()

    def _workbook_bytes(self, configure=None) -> bytes:
        stream = io.BytesIO()
        wb = Workbook()
        wb.active.title = "Sheet1"
        if configure is not None:
            configure(wb)
        wb.save(stream)
        wb.close()
        return stream.getvalue()

    def test_validate_rejects_http_scheme(self):
        with self.assertRaises(remote.RemoteWorkbookError):
            remote.validate_signed_url(
                f"http://{ALLOWED_HOST}/a.xlsx", purpose="input download"
            )

    def test_validate_rejects_unallowed_host(self):
        with self.assertRaises(remote.RemoteWorkbookError):
            remote.validate_signed_url(
                "https://evil.example.com/a.xlsx", purpose="input download"
            )

    def test_validate_rejects_embedded_credentials(self):
        with self.assertRaises(remote.RemoteWorkbookError):
            remote.validate_signed_url(
                f"https://user:pass@{ALLOWED_HOST}/a.xlsx",
                purpose="input download",
            )

    def test_create_uploads_workbook(self):
        put_response = MagicMock(status_code=200, is_redirect=False)
        fake_client = _FakeClient(put_response=put_response)

        with patch.object(remote, "_create_http_client", return_value=fake_client):
            result = remote.execute_workbook_job(
                operation="create",
                request_id="req-create",
                output_upload_url=_url("/out.xlsx"),
                upload_headers={
                    "Content-Type": remote.XLSX_CONTENT_TYPE,
                },
            )

        self.assertTrue(result["success"])
        self.assertEqual(result["operation"], "create")
        self.assertTrue(result["output_uploaded"])
        self.assertGreater(result["size_bytes"], 0)
        self.assertEqual(len(fake_client.put_calls), 1)
        self.assertEqual(
            fake_client.put_calls[0]["headers"]["Content-Type"],
            remote.XLSX_CONTENT_TYPE,
        )
        # Response body should be a valid xlsx.
        wb = load_workbook(io.BytesIO(fake_client.put_calls[0]["body"]))
        self.assertIn("Sheet1", wb.sheetnames)
        wb.close()

    def test_write_downloads_writes_and_uploads(self):
        source = self._xlsx_bytes()
        get_response = _FakeStreamResponse(
            source, headers={"Content-Length": str(len(source))}
        )
        put_response = MagicMock(status_code=200, is_redirect=False)
        fake_client = _FakeClient(
            get_response=get_response, put_response=put_response
        )

        with patch.object(remote, "_create_http_client", return_value=fake_client):
            result = remote.execute_workbook_job(
                operation="write",
                request_id="req-write",
                input_download_url=_url("/in.xlsx"),
                output_upload_url=_url("/out.xlsx"),
                upload_headers={"Content-Type": remote.XLSX_CONTENT_TYPE},
                sheet_name="Sheet1",
                start_cell="A1",
                data=[["Name", "Score"], ["Nakul", 100]],
            )

        self.assertTrue(result["success"])
        self.assertEqual(result["operation"], "write")
        self.assertTrue(result["output_uploaded"])
        self.assertEqual(len(fake_client.put_calls), 1)

        wb = load_workbook(io.BytesIO(fake_client.put_calls[0]["body"]))
        ws = wb["Sheet1"]
        self.assertEqual(ws["A1"].value, "Name")
        self.assertEqual(ws["B2"].value, 100)
        wb.close()

    def test_read_preserves_blank_rows(self):
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "blank.xlsx")
            wb = Workbook()
            ws = wb.active
            ws.title = "Sheet1"
            ws["A1"] = "Name"
            ws["B1"] = "Score"
            ws["C1"] = "Status"
            ws["A2"] = "Nakul"
            ws["B2"] = 100
            ws["C2"] = "Active"
            # Row 3 left blank intentionally.
            ws["A4"] = "Sara"
            ws["B4"] = 95
            ws["C4"] = "Active"
            wb.save(path)
            wb.close()
            with open(path, "rb") as f:
                source = f.read()

        get_response = _FakeStreamResponse(source)
        fake_client = _FakeClient(get_response=get_response)

        with patch.object(remote, "_create_http_client", return_value=fake_client):
            result = remote.execute_workbook_job(
                operation="read",
                request_id="req-read",
                input_download_url=_url("/in.xlsx"),
                sheet_name="Sheet1",
                start_cell="A1",
                end_cell="C4",
            )

        self.assertTrue(result["success"])
        self.assertEqual(result["operation"], "read")
        self.assertFalse(result["output_uploaded"])
        values = result["data"]["values"]
        self.assertEqual(len(values), 4)
        self.assertEqual(values[0], ["Name", "Score", "Status"])
        self.assertEqual(values[1], ["Nakul", 100, "Active"])
        self.assertEqual(values[2], [None, None, None])
        self.assertEqual(values[3], ["Sara", 95, "Active"])

    def test_preview_only_limits_to_ten_rows(self):
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "many.xlsx")
            wb = Workbook()
            ws = wb.active
            ws.title = "Sheet1"
            for i in range(1, 16):
                ws.cell(row=i, column=1, value=i)
            wb.save(path)
            wb.close()
            with open(path, "rb") as f:
                source = f.read()

        get_response = _FakeStreamResponse(source)
        fake_client = _FakeClient(get_response=get_response)

        with patch.object(remote, "_create_http_client", return_value=fake_client):
            result = remote.execute_workbook_job(
                operation="read",
                input_download_url=_url("/in.xlsx"),
                sheet_name="Sheet1",
                start_cell="A1",
                end_cell="A15",
                preview_only=True,
            )

        self.assertTrue(result["success"])
        self.assertEqual(len(result["data"]["values"]), 10)
        self.assertEqual(result["data"]["end_cell"], "A10")

    def test_download_rejects_declared_oversize(self):
        get_response = _FakeStreamResponse(
            b"abc",
            headers={"Content-Length": str(2 * 1024 * 1024)},
        )
        fake_client = _FakeClient(get_response=get_response)

        with patch.object(remote, "_create_http_client", return_value=fake_client):
            result = remote.execute_workbook_job(
                operation="read",
                input_download_url=_url("/big.xlsx"),
                sheet_name="Sheet1",
            )

        self.assertFalse(result["success"])
        self.assertIn("maximum size", result["error"])

    def test_download_rejects_streamed_oversize(self):
        os.environ["EXCEL_MCP_MAX_FILE_BYTES"] = "20"
        get_response = _FakeStreamResponse(b"x" * 40)
        fake_client = _FakeClient(get_response=get_response)

        with patch.object(remote, "_create_http_client", return_value=fake_client):
            result = remote.execute_workbook_job(
                operation="read",
                input_download_url=_url("/big.xlsx"),
                sheet_name="Sheet1",
            )

        self.assertFalse(result["success"])
        self.assertIn("maximum size", result["error"])

    def test_upload_non_2xx_is_failure(self):
        put_response = MagicMock(status_code=403, is_redirect=False)
        fake_client = _FakeClient(put_response=put_response)

        with patch.object(remote, "_create_http_client", return_value=fake_client):
            result = remote.execute_workbook_job(
                operation="create",
                output_upload_url=_url("/out.xlsx"),
                upload_headers={"Content-Type": remote.XLSX_CONTENT_TYPE},
            )

        self.assertFalse(result["success"])
        self.assertIn("403", result["error"])

    def test_temp_dir_cleaned_after_success(self):
        put_response = MagicMock(status_code=200, is_redirect=False)
        fake_client = _FakeClient(put_response=put_response)
        created = {}

        real_mkdtemp = tempfile.mkdtemp

        def tracking_mkdtemp(*args, **kwargs):
            path = real_mkdtemp(*args, **kwargs)
            created["path"] = path
            return path

        with patch.object(tempfile, "mkdtemp", side_effect=tracking_mkdtemp):
            with patch.object(
                remote, "_create_http_client", return_value=fake_client
            ):
                result = remote.execute_workbook_job(
                    operation="create",
                    output_upload_url=_url("/out.xlsx"),
                    upload_headers={"Content-Type": remote.XLSX_CONTENT_TYPE},
                )

        self.assertTrue(result["success"])
        self.assertIn("path", created)
        self.assertFalse(os.path.exists(created["path"]))

    def test_temp_dir_cleaned_after_failure(self):
        put_response = MagicMock(status_code=500, is_redirect=False)
        fake_client = _FakeClient(put_response=put_response)
        created = {}
        real_mkdtemp = tempfile.mkdtemp

        def tracking_mkdtemp(*args, **kwargs):
            path = real_mkdtemp(*args, **kwargs)
            created["path"] = path
            return path

        with patch.object(tempfile, "mkdtemp", side_effect=tracking_mkdtemp):
            with patch.object(
                remote, "_create_http_client", return_value=fake_client
            ):
                result = remote.execute_workbook_job(
                    operation="create",
                    output_upload_url=_url("/out.xlsx"),
                    upload_headers={"Content-Type": remote.XLSX_CONTENT_TYPE},
                )

        self.assertFalse(result["success"])
        self.assertFalse(os.path.exists(created["path"]))

    def test_unknown_operation_rejected(self):
        result = remote.execute_workbook_job(operation="format")
        self.assertFalse(result["success"])
        self.assertIn("Unsupported operation", result["error"])

    def test_error_messages_do_not_include_urls(self):
        result = remote.execute_workbook_job(
            operation="create",
            output_upload_url="https://evil.example.com/secret-token",
            upload_headers={"Content-Type": remote.XLSX_CONTENT_TYPE},
        )
        self.assertFalse(result["success"])
        self.assertNotIn("evil.example.com", result["error"])
        self.assertNotIn("secret-token", result["error"])

    def test_apply_formula_normalizes_and_uploads(self):
        source = self._workbook_bytes()
        fake_client = _FakeClient(
            get_response=_FakeStreamResponse(source),
            put_response=MagicMock(status_code=200, is_redirect=False),
        )

        with patch.object(remote, "_create_http_client", return_value=fake_client):
            result = remote.execute_workbook_job(
                operation="apply_formula",
                input_download_url=_url("/in.xlsx"),
                output_upload_url=_url("/out.xlsx"),
                sheet_name="Sheet1",
                cell="B2",
                formula="SUM(A1:A10)",
            )

        self.assertTrue(result["success"])
        self.assertTrue(result["output_uploaded"])
        wb = load_workbook(io.BytesIO(fake_client.put_calls[0]["body"]))
        self.assertEqual(wb["Sheet1"]["B2"].value, "=SUM(A1:A10)")
        wb.close()

    def test_validate_formula_returns_match_without_put(self):
        def configure(wb):
            wb["Sheet1"]["B2"] = "=SUM(A1:A10)"

        source = self._workbook_bytes(configure)
        fake_client = _FakeClient(
            get_response=_FakeStreamResponse(source),
        )

        with patch.object(remote, "_create_http_client", return_value=fake_client):
            result = remote.execute_workbook_job(
                operation="validate_formula",
                input_download_url=_url("/in.xlsx"),
                sheet_name="Sheet1",
                cell="B2",
                formula="=SUM(A1:A10)",
            )

        self.assertTrue(result["success"])
        self.assertFalse(result["output_uploaded"])
        self.assertTrue(result["data"]["valid"])
        self.assertTrue(result["data"]["matches"])
        self.assertEqual(result["data"]["current_formula"], "=SUM(A1:A10)")
        self.assertEqual(fake_client.put_calls, [])

    def test_validate_formula_rejects_blocked_function(self):
        source = self._workbook_bytes()
        fake_client = _FakeClient(
            get_response=_FakeStreamResponse(source),
        )

        with patch.object(remote, "_create_http_client", return_value=fake_client):
            result = remote.execute_workbook_job(
                operation="validate_formula",
                input_download_url=_url("/in.xlsx"),
                sheet_name="Sheet1",
                cell="A1",
                formula="=WEBSERVICE(\"https://example.com\")",
            )

        self.assertFalse(result["success"])
        self.assertIn("Unsafe function", result["error"])
        self.assertEqual(fake_client.put_calls, [])

    def test_validate_formula_requires_equals_prefix(self):
        source = self._workbook_bytes()
        fake_client = _FakeClient(
            get_response=_FakeStreamResponse(source),
        )

        with patch.object(remote, "_create_http_client", return_value=fake_client):
            result = remote.execute_workbook_job(
                operation="validate_formula",
                input_download_url=_url("/in.xlsx"),
                sheet_name="Sheet1",
                cell="A1",
                formula="SUM(A1:A2)",
            )

        self.assertFalse(result["success"])
        self.assertIn("must start with '='", result["error"])

    def test_create_worksheet_adds_sheet(self):
        source = self._workbook_bytes()
        fake_client = _FakeClient(
            get_response=_FakeStreamResponse(source),
            put_response=MagicMock(status_code=200, is_redirect=False),
        )

        with patch.object(remote, "_create_http_client", return_value=fake_client):
            result = remote.execute_workbook_job(
                operation="create_worksheet",
                input_download_url=_url("/in.xlsx"),
                output_upload_url=_url("/out.xlsx"),
                sheet_name="Summary",
            )

        self.assertTrue(result["success"])
        wb = load_workbook(io.BytesIO(fake_client.put_calls[0]["body"]))
        self.assertIn("Summary", wb.sheetnames)
        wb.close()

    def test_copy_worksheet_preserves_cell_data(self):
        def configure(wb):
            wb["Sheet1"]["A1"] = "copied value"

        source = self._workbook_bytes(configure)
        fake_client = _FakeClient(
            get_response=_FakeStreamResponse(source),
            put_response=MagicMock(status_code=200, is_redirect=False),
        )

        with patch.object(remote, "_create_http_client", return_value=fake_client):
            result = remote.execute_workbook_job(
                operation="copy_worksheet",
                input_download_url=_url("/in.xlsx"),
                output_upload_url=_url("/out.xlsx"),
                source_sheet="Sheet1",
                target_sheet="Sheet1 Copy",
            )

        self.assertTrue(result["success"])
        wb = load_workbook(io.BytesIO(fake_client.put_calls[0]["body"]))
        self.assertEqual(wb["Sheet1 Copy"]["A1"].value, "copied value")
        wb.close()

    def test_rename_worksheet_changes_name(self):
        source = self._workbook_bytes()
        fake_client = _FakeClient(
            get_response=_FakeStreamResponse(source),
            put_response=MagicMock(status_code=200, is_redirect=False),
        )

        with patch.object(remote, "_create_http_client", return_value=fake_client):
            result = remote.execute_workbook_job(
                operation="rename_worksheet",
                input_download_url=_url("/in.xlsx"),
                output_upload_url=_url("/out.xlsx"),
                old_name="Sheet1",
                new_name="January",
            )

        self.assertTrue(result["success"])
        wb = load_workbook(io.BytesIO(fake_client.put_calls[0]["body"]))
        self.assertIn("January", wb.sheetnames)
        self.assertNotIn("Sheet1", wb.sheetnames)
        wb.close()

    def test_delete_worksheet_removes_sheet(self):
        def configure(wb):
            wb.create_sheet("Draft")

        source = self._workbook_bytes(configure)
        fake_client = _FakeClient(
            get_response=_FakeStreamResponse(source),
            put_response=MagicMock(status_code=200, is_redirect=False),
        )

        with patch.object(remote, "_create_http_client", return_value=fake_client):
            result = remote.execute_workbook_job(
                operation="delete_worksheet",
                input_download_url=_url("/in.xlsx"),
                output_upload_url=_url("/out.xlsx"),
                sheet_name="Draft",
            )

        self.assertTrue(result["success"])
        wb = load_workbook(io.BytesIO(fake_client.put_calls[0]["body"]))
        self.assertNotIn("Draft", wb.sheetnames)
        wb.close()

    def test_delete_only_worksheet_fails_without_put(self):
        source = self._workbook_bytes()
        fake_client = _FakeClient(
            get_response=_FakeStreamResponse(source),
        )

        with patch.object(remote, "_create_http_client", return_value=fake_client):
            result = remote.execute_workbook_job(
                operation="delete_worksheet",
                input_download_url=_url("/in.xlsx"),
                output_upload_url=_url("/out.xlsx"),
                sheet_name="Sheet1",
            )

        self.assertFalse(result["success"])
        self.assertIn("only sheet", result["error"])
        self.assertEqual(fake_client.put_calls, [])

    def test_invalid_worksheet_name_fails_before_download(self):
        fake_client = _FakeClient()

        with patch.object(remote, "_create_http_client", return_value=fake_client):
            result = remote.execute_workbook_job(
                operation="create_worksheet",
                input_download_url=_url("/in.xlsx"),
                output_upload_url=_url("/out.xlsx"),
                sheet_name="Bad:Name",
            )

        self.assertFalse(result["success"])
        self.assertIn("invalid character", result["error"])
        self.assertEqual(fake_client.get_calls, [])

    def test_duplicate_worksheet_name_fails_without_put(self):
        source = self._workbook_bytes()
        fake_client = _FakeClient(
            get_response=_FakeStreamResponse(source),
        )

        with patch.object(remote, "_create_http_client", return_value=fake_client):
            result = remote.execute_workbook_job(
                operation="create_worksheet",
                input_download_url=_url("/in.xlsx"),
                output_upload_url=_url("/out.xlsx"),
                sheet_name="Sheet1",
            )

        self.assertFalse(result["success"])
        self.assertIn("already exists", result["error"])
        self.assertEqual(fake_client.put_calls, [])

    def test_missing_source_worksheet_fails_without_put(self):
        source = self._workbook_bytes()
        fake_client = _FakeClient(
            get_response=_FakeStreamResponse(source),
        )

        with patch.object(remote, "_create_http_client", return_value=fake_client):
            result = remote.execute_workbook_job(
                operation="copy_worksheet",
                input_download_url=_url("/in.xlsx"),
                output_upload_url=_url("/out.xlsx"),
                source_sheet="Missing",
                target_sheet="Copy",
            )

        self.assertFalse(result["success"])
        self.assertIn("not found", result["error"])
        self.assertEqual(fake_client.put_calls, [])

    def test_overlong_worksheet_name_fails_before_download(self):
        fake_client = _FakeClient()

        with patch.object(remote, "_create_http_client", return_value=fake_client):
            result = remote.execute_workbook_job(
                operation="create_worksheet",
                input_download_url=_url("/in.xlsx"),
                output_upload_url=_url("/out.xlsx"),
                sheet_name="x" * 32,
            )

        self.assertFalse(result["success"])
        self.assertIn("31 characters or fewer", result["error"])
        self.assertEqual(fake_client.get_calls, [])

    def test_missing_phase_two_field_fails_before_download(self):
        fake_client = _FakeClient()

        with patch.object(remote, "_create_http_client", return_value=fake_client):
            result = remote.execute_workbook_job(
                operation="copy_worksheet",
                input_download_url=_url("/in.xlsx"),
                output_upload_url=_url("/out.xlsx"),
                source_sheet="Sheet1",
            )

        self.assertFalse(result["success"])
        self.assertIn("target_sheet is required", result["error"])
        self.assertEqual(fake_client.get_calls, [])

    def test_format_range_preserves_unrelated_font(self):
        def configure(wb):
            cell = wb["Sheet1"]["A1"]
            cell.value = "Header"
            cell.font = Font(bold=True, name="Calibri", size=14)

        source = self._workbook_bytes(configure)
        fake_client = _FakeClient(
            get_response=_FakeStreamResponse(source),
            put_response=MagicMock(status_code=200, is_redirect=False),
        )

        with patch.object(remote, "_create_http_client", return_value=fake_client):
            result = remote.execute_workbook_job(
                operation="format_range",
                input_download_url=_url("/in.xlsx"),
                output_upload_url=_url("/out.xlsx"),
                sheet_name="Sheet1",
                start_cell="A1",
                format_options={"bg_color": "217346"},
            )

        self.assertTrue(result["success"])
        wb = load_workbook(io.BytesIO(fake_client.put_calls[0]["body"]))
        cell = wb["Sheet1"]["A1"]
        self.assertTrue(cell.font.bold)
        self.assertEqual(cell.font.size, 14)
        self.assertEqual(cell.fill.fgColor.rgb, "FF217346")
        wb.close()

    def test_format_range_rejects_empty_options(self):
        fake_client = _FakeClient()
        with patch.object(remote, "_create_http_client", return_value=fake_client):
            result = remote.execute_workbook_job(
                operation="format_range",
                input_download_url=_url("/in.xlsx"),
                output_upload_url=_url("/out.xlsx"),
                sheet_name="Sheet1",
                start_cell="A1",
                format_options={},
            )
        self.assertFalse(result["success"])
        self.assertIn("non-empty", result["error"])
        self.assertEqual(fake_client.get_calls, [])

    def test_format_range_rejects_unknown_key_and_bad_color(self):
        fake_client = _FakeClient()
        with patch.object(remote, "_create_http_client", return_value=fake_client):
            result = remote.execute_workbook_job(
                operation="format_range",
                input_download_url=_url("/in.xlsx"),
                output_upload_url=_url("/out.xlsx"),
                sheet_name="Sheet1",
                start_cell="A1",
                format_options={"merge_cells": True},
            )
        self.assertFalse(result["success"])
        self.assertIn("Unknown format_options", result["error"])

        with patch.object(remote, "_create_http_client", return_value=fake_client):
            result = remote.execute_workbook_job(
                operation="format_range",
                input_download_url=_url("/in.xlsx"),
                output_upload_url=_url("/out.xlsx"),
                sheet_name="Sheet1",
                start_cell="A1",
                format_options={"font_color": "#FFFFFF"},
            )
        self.assertFalse(result["success"])
        self.assertIn("#", result["error"])

    def test_merge_and_unmerge_cells(self):
        source = self._workbook_bytes()
        fake_client = _FakeClient(
            get_response=_FakeStreamResponse(source),
            put_response=MagicMock(status_code=200, is_redirect=False),
        )
        with patch.object(remote, "_create_http_client", return_value=fake_client):
            result = remote.execute_workbook_job(
                operation="merge_cells",
                input_download_url=_url("/in.xlsx"),
                output_upload_url=_url("/out.xlsx"),
                sheet_name="Sheet1",
                start_cell="A1",
                end_cell="C1",
            )
        self.assertTrue(result["success"])
        wb = load_workbook(io.BytesIO(fake_client.put_calls[0]["body"]))
        self.assertIn("A1:C1", [str(r) for r in wb["Sheet1"].merged_cells.ranges])
        merged_bytes = fake_client.put_calls[0]["body"]
        wb.close()

        fake_client2 = _FakeClient(
            get_response=_FakeStreamResponse(merged_bytes),
            put_response=MagicMock(status_code=200, is_redirect=False),
        )
        with patch.object(remote, "_create_http_client", return_value=fake_client2):
            result = remote.execute_workbook_job(
                operation="unmerge_cells",
                input_download_url=_url("/in.xlsx"),
                output_upload_url=_url("/out.xlsx"),
                sheet_name="Sheet1",
                start_cell="A1",
                end_cell="C1",
            )
        self.assertTrue(result["success"])

    def test_reject_overlapping_merge(self):
        def configure(wb):
            wb["Sheet1"].merge_cells("A1:C1")

        source = self._workbook_bytes(configure)
        fake_client = _FakeClient(get_response=_FakeStreamResponse(source))
        with patch.object(remote, "_create_http_client", return_value=fake_client):
            result = remote.execute_workbook_job(
                operation="merge_cells",
                input_download_url=_url("/in.xlsx"),
                output_upload_url=_url("/out.xlsx"),
                sheet_name="Sheet1",
                start_cell="B1",
                end_cell="D1",
            )
        self.assertFalse(result["success"])
        self.assertIn("overlaps", result["error"])
        self.assertEqual(fake_client.put_calls, [])

    def test_copy_range_overlapping_same_sheet(self):
        def configure(wb):
            wb["Sheet1"]["A1"] = "one"
            wb["Sheet1"]["A2"] = "two"
            wb["Sheet1"]["A3"] = "three"

        source = self._workbook_bytes(configure)
        fake_client = _FakeClient(
            get_response=_FakeStreamResponse(source),
            put_response=MagicMock(status_code=200, is_redirect=False),
        )
        with patch.object(remote, "_create_http_client", return_value=fake_client):
            result = remote.execute_workbook_job(
                operation="copy_range",
                input_download_url=_url("/in.xlsx"),
                output_upload_url=_url("/out.xlsx"),
                sheet_name="Sheet1",
                source_start="A1",
                source_end="A3",
                target_start="A2",
            )
        self.assertTrue(result["success"])
        wb = load_workbook(io.BytesIO(fake_client.put_calls[0]["body"]))
        self.assertEqual(wb["Sheet1"]["A2"].value, "one")
        self.assertEqual(wb["Sheet1"]["A3"].value, "two")
        self.assertEqual(wb["Sheet1"]["A4"].value, "three")
        wb.close()

    def test_delete_range_shifts_up_within_columns(self):
        def configure(wb):
            ws = wb["Sheet1"]
            ws["A1"] = "keep-a"
            ws["B1"] = "del-1"
            ws["C1"] = "keep-c"
            ws["B2"] = "del-2"
            ws["B3"] = "move-up"
            ws["A3"] = "stay-a"
            ws["C3"] = "stay-c"

        source = self._workbook_bytes(configure)
        fake_client = _FakeClient(
            get_response=_FakeStreamResponse(source),
            put_response=MagicMock(status_code=200, is_redirect=False),
        )
        with patch.object(remote, "_create_http_client", return_value=fake_client):
            result = remote.execute_workbook_job(
                operation="delete_range",
                input_download_url=_url("/in.xlsx"),
                output_upload_url=_url("/out.xlsx"),
                sheet_name="Sheet1",
                start_cell="B1",
                end_cell="B2",
                shift_direction="up",
            )
        self.assertTrue(result["success"])
        wb = load_workbook(io.BytesIO(fake_client.put_calls[0]["body"]))
        ws = wb["Sheet1"]
        self.assertEqual(ws["A1"].value, "keep-a")
        self.assertEqual(ws["C1"].value, "keep-c")
        self.assertEqual(ws["B1"].value, "move-up")
        self.assertEqual(ws["A3"].value, "stay-a")
        self.assertEqual(ws["C3"].value, "stay-c")
        wb.close()

    def test_get_merged_cells_and_metadata_no_put(self):
        def configure(wb):
            wb["Sheet1"]["A1"] = "x"
            wb["Sheet1"].merge_cells("A1:B1")
            wb.create_sheet("Empty")

        source = self._workbook_bytes(configure)
        fake_client = _FakeClient(get_response=_FakeStreamResponse(source))
        with patch.object(remote, "_create_http_client", return_value=fake_client):
            merged = remote.execute_workbook_job(
                operation="get_merged_cells",
                input_download_url=_url("/in.xlsx"),
                sheet_name="Sheet1",
            )
            meta = remote.execute_workbook_job(
                operation="get_workbook_metadata",
                input_download_url=_url("/in.xlsx"),
                include_ranges=True,
            )
        self.assertTrue(merged["success"])
        self.assertEqual(merged["data"]["merged_ranges"], ["A1:B1"])
        self.assertTrue(meta["success"])
        self.assertIn("Sheet1", meta["data"]["sheets"])
        self.assertIn("size_bytes", meta["data"])
        self.assertNotIn("filename", meta["data"])
        self.assertNotIn("modified_at", meta["data"])
        self.assertIn("Sheet1", meta["data"]["used_ranges"])
        self.assertNotIn("Empty", meta["data"]["used_ranges"])
        self.assertEqual(fake_client.put_calls, [])

    def test_validate_excel_range_structured(self):
        def configure(wb):
            wb["Sheet1"]["A1"] = "a"
            wb["Sheet1"]["B2"] = "b"

        source = self._workbook_bytes(configure)
        fake_client = _FakeClient(get_response=_FakeStreamResponse(source))
        with patch.object(remote, "_create_http_client", return_value=fake_client):
            result = remote.execute_workbook_job(
                operation="validate_excel_range",
                input_download_url=_url("/in.xlsx"),
                sheet_name="Sheet1",
                start_cell="A1",
                end_cell="B2",
            )
        self.assertTrue(result["success"])
        self.assertTrue(result["data"]["valid"])
        self.assertEqual(result["data"]["range"], "A1:B2")
        self.assertEqual(fake_client.put_calls, [])

    def test_rejects_malformed_cell_reference(self):
        fake_client = _FakeClient()
        with patch.object(remote, "_create_http_client", return_value=fake_client):
            result = remote.execute_workbook_job(
                operation="merge_cells",
                input_download_url=_url("/in.xlsx"),
                output_upload_url=_url("/out.xlsx"),
                sheet_name="Sheet1",
                start_cell="A1junk",
                end_cell="B2",
            )
        self.assertFalse(result["success"])
        self.assertIn("Invalid cell reference", result["error"])
        self.assertEqual(fake_client.get_calls, [])

    def test_get_data_validation_info_empty(self):
        def configure(wb):
            wb["Sheet1"]["A1"] = "x"

        source = self._workbook_bytes(configure)
        fake_client = _FakeClient(get_response=_FakeStreamResponse(source))
        with patch.object(remote, "_create_http_client", return_value=fake_client):
            result = remote.execute_workbook_job(
                operation="get_data_validation_info",
                input_download_url=_url("/in.xlsx"),
                sheet_name="Sheet1",
            )
        self.assertTrue(result["success"])
        self.assertEqual(result["operation"], "get_data_validation_info")
        self.assertFalse(result["output_uploaded"])
        self.assertEqual(
            result["data"],
            {"sheet_name": "Sheet1", "validation_rules": []},
        )
        self.assertEqual(fake_client.put_calls, [])

    def test_get_data_validation_info_with_list(self):
        from openpyxl.worksheet.datavalidation import DataValidation

        def configure(wb):
            ws = wb["Sheet1"]
            ws["A1"] = "Status"
            ws["A2"] = "Open"
            dv = DataValidation(
                type="list",
                formula1='"Open,Closed"',
                allow_blank=True,
            )
            dv.add("B2:B20")
            ws.add_data_validation(dv)

        source = self._workbook_bytes(configure)
        fake_client = _FakeClient(get_response=_FakeStreamResponse(source))
        with patch.object(remote, "_create_http_client", return_value=fake_client):
            result = remote.execute_workbook_job(
                operation="get_data_validation_info",
                input_download_url=_url("/in.xlsx"),
                sheet_name="Sheet1",
            )
        self.assertTrue(result["success"])
        rules = result["data"]["validation_rules"]
        self.assertEqual(len(rules), 1)
        self.assertEqual(rules[0]["validation_type"], "list")
        self.assertEqual(rules[0]["ranges"], "B2:B20")
        self.assertEqual(rules[0]["allowed_values"], ["Open", "Closed"])

    def test_create_chart_uploads_and_places_chart(self):
        def configure(wb):
            ws = wb["Sheet1"]
            ws["A1"] = "Name"
            ws["B1"] = "Score"
            ws["A2"] = "Ada"
            ws["B2"] = 10
            ws["A3"] = "Bob"
            ws["B3"] = 20

        source = self._workbook_bytes(configure)
        put_response = MagicMock(status_code=200, is_redirect=False)
        fake_client = _FakeClient(
            get_response=_FakeStreamResponse(source),
            put_response=put_response,
        )
        with patch.object(remote, "_create_http_client", return_value=fake_client):
            result = remote.execute_workbook_job(
                operation="create_chart",
                input_download_url=_url("/in.xlsx"),
                output_upload_url=_url("/out.xlsx"),
                sheet_name="Sheet1",
                data_range="A1:B3",
                chart_type="bar",
                target_cell="AA1",
                title="Scores",
            )
        self.assertTrue(result["success"])
        self.assertTrue(result["output_uploaded"])
        self.assertIsNone(result["data"])
        wb = load_workbook(io.BytesIO(fake_client.put_calls[0]["body"]))
        try:
            self.assertEqual(len(wb["Sheet1"]._charts), 1)
        finally:
            wb.close()

    def test_create_chart_rejects_unsupported_type(self):
        fake_client = _FakeClient()
        with patch.object(remote, "_create_http_client", return_value=fake_client):
            result = remote.execute_workbook_job(
                operation="create_chart",
                input_download_url=_url("/in.xlsx"),
                output_upload_url=_url("/out.xlsx"),
                sheet_name="Sheet1",
                data_range="A1:B3",
                chart_type="bubble",
                target_cell="E2",
            )
        self.assertFalse(result["success"])
        self.assertIn("chart_type", result["error"])
        self.assertEqual(fake_client.get_calls, [])

    def test_create_table_uploads(self):
        def configure(wb):
            ws = wb["Sheet1"]
            ws["A1"] = "Name"
            ws["B1"] = "Score"
            ws["A2"] = "Ada"
            ws["B2"] = 10

        source = self._workbook_bytes(configure)
        put_response = MagicMock(status_code=200, is_redirect=False)
        fake_client = _FakeClient(
            get_response=_FakeStreamResponse(source),
            put_response=put_response,
        )
        with patch.object(remote, "_create_http_client", return_value=fake_client):
            result = remote.execute_workbook_job(
                operation="create_table",
                input_download_url=_url("/in.xlsx"),
                output_upload_url=_url("/out.xlsx"),
                sheet_name="Sheet1",
                data_range="A1:B2",
                table_name="ScoresTable",
            )
        self.assertTrue(result["success"])
        wb = load_workbook(io.BytesIO(fake_client.put_calls[0]["body"]))
        try:
            self.assertIn("ScoresTable", wb["Sheet1"].tables)
        finally:
            wb.close()

    def test_create_table_rejects_bad_style(self):
        def configure(wb):
            ws = wb["Sheet1"]
            ws["A1"] = "Name"
            ws["B1"] = "Score"
            ws["A2"] = "Ada"
            ws["B2"] = 10

        source = self._workbook_bytes(configure)
        fake_client = _FakeClient(get_response=_FakeStreamResponse(source))
        with patch.object(remote, "_create_http_client", return_value=fake_client):
            result = remote.execute_workbook_job(
                operation="create_table",
                input_download_url=_url("/in.xlsx"),
                output_upload_url=_url("/out.xlsx"),
                sheet_name="Sheet1",
                data_range="A1:B2",
                table_style="NotARealStyle",
            )
        self.assertFalse(result["success"])
        self.assertIn("table_style", result["error"])
        self.assertEqual(fake_client.put_calls, [])

    def test_create_pivot_table_with_column_grouping(self):
        def configure(wb):
            ws = wb["Sheet1"]
            headers = ["Region", "Category", "Score"]
            rows = [
                ["East", "A", 10],
                ["East", "B", 5],
                ["West", "A", 7],
                ["West", "B", "3"],  # numeric string
            ]
            for col, header in enumerate(headers, start=1):
                ws.cell(1, col, header)
            for r_idx, row in enumerate(rows, start=2):
                for c_idx, value in enumerate(row, start=1):
                    ws.cell(r_idx, c_idx, value)

        source = self._workbook_bytes(configure)
        put_response = MagicMock(status_code=200, is_redirect=False)
        fake_client = _FakeClient(
            get_response=_FakeStreamResponse(source),
            put_response=put_response,
        )
        with patch.object(remote, "_create_http_client", return_value=fake_client):
            result = remote.execute_workbook_job(
                operation="create_pivot_table",
                input_download_url=_url("/in.xlsx"),
                output_upload_url=_url("/out.xlsx"),
                sheet_name="Sheet1",
                data_range="A1:C5",
                rows=["Region"],
                columns=["Category"],
                values=["Score"],
                agg_func="sum",
            )
        self.assertTrue(result["success"])
        self.assertIsNone(result["data"])
        wb = load_workbook(io.BytesIO(fake_client.put_calls[0]["body"]))
        try:
            self.assertIn("Sheet1_pivot", wb.sheetnames)
            pivot = wb["Sheet1_pivot"]
            headers = [pivot.cell(1, c).value for c in range(1, 4)]
            self.assertEqual(headers[0], "Region")
            self.assertTrue(any("Category=A" in str(h) for h in headers[1:]))
            self.assertTrue(any("Category=B" in str(h) for h in headers[1:]))
            # Find East row and verify A=10, B=5
            east_row = None
            for r in range(2, pivot.max_row + 1):
                if pivot.cell(r, 1).value == "East":
                    east_row = r
                    break
            self.assertIsNotNone(east_row)
            values_by_header = {
                pivot.cell(1, c).value: pivot.cell(east_row, c).value
                for c in range(2, 4)
            }
            a_key = next(k for k in values_by_header if "Category=A" in str(k))
            b_key = next(k for k in values_by_header if "Category=B" in str(k))
            self.assertEqual(values_by_header[a_key], 10)
            self.assertEqual(values_by_header[b_key], 5)
        finally:
            wb.close()

    def test_create_pivot_rejects_mean(self):
        fake_client = _FakeClient()
        with patch.object(remote, "_create_http_client", return_value=fake_client):
            result = remote.execute_workbook_job(
                operation="create_pivot_table",
                input_download_url=_url("/in.xlsx"),
                output_upload_url=_url("/out.xlsx"),
                sheet_name="Sheet1",
                data_range="A1:C5",
                rows=["Region"],
                values=["Score"],
                agg_func="mean",
            )
        self.assertFalse(result["success"])
        self.assertIn("agg_func", result["error"])
        self.assertEqual(fake_client.get_calls, [])


class Phase5LocalHelperTests(unittest.TestCase):
    def test_chart_uses_add_chart_and_supports_aa1(self):
        from excel_mcp.chart import create_chart_in_sheet

        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "chart.xlsx")
            wb = Workbook()
            ws = wb.active
            ws.title = "Sheet1"
            ws["A1"] = "Name"
            ws["B1"] = "Score"
            ws["A2"] = "Ada"
            ws["B2"] = 10
            wb.save(path)
            wb.close()

            create_chart_in_sheet(
                path,
                "Sheet1",
                "A1:B2",
                "line",
                "AA1",
                title="T",
            )
            reopened = load_workbook(path)
            try:
                self.assertEqual(len(reopened["Sheet1"]._charts), 1)
            finally:
                reopened.close()

    def test_table_rejects_cell_like_name(self):
        from excel_mcp.tables import create_excel_table
        from excel_mcp.exceptions import DataError

        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "table.xlsx")
            wb = Workbook()
            ws = wb.active
            ws.title = "Sheet1"
            ws["A1"] = "Name"
            ws["B1"] = "Score"
            ws["A2"] = "Ada"
            ws["B2"] = 10
            wb.save(path)
            wb.close()

            with self.assertRaises(DataError):
                create_excel_table(
                    path, "Sheet1", "A1:B2", table_name="A1"
                )

    def test_pivot_numeric_string_grouping_and_replace(self):
        from excel_mcp.pivot import create_pivot_table

        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "pivot.xlsx")
            wb = Workbook()
            ws = wb.active
            ws.title = "Sheet1"
            for col, header in enumerate(["Region", "Score"], start=1):
                ws.cell(1, col, header)
            ws.cell(2, 1, 10)
            ws.cell(2, 2, 5)
            ws.cell(3, 1, "10")
            ws.cell(3, 2, 7)
            wb.save(path)
            wb.close()

            create_pivot_table(
                path,
                "Sheet1",
                "A1:B3",
                rows=["Region"],
                values=["Score"],
                agg_func="sum",
            )
            create_pivot_table(
                path,
                "Sheet1",
                "A1:B3",
                rows=["Region"],
                values=["Score"],
                agg_func="sum",
            )
            reopened = load_workbook(path)
            try:
                self.assertEqual(
                    reopened.sheetnames.count("Sheet1_pivot"), 1
                )
                pivot = reopened["Sheet1_pivot"]
                # One grouped row for key "10" with sum 12
                self.assertEqual(pivot["A2"].value, "10")
                self.assertEqual(pivot["B2"].value, 12)
            finally:
                reopened.close()

    def test_validation_info_always_structured(self):
        from excel_mcp.cell_validation import get_data_validation_info

        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "val.xlsx")
            wb = Workbook()
            wb.active.title = "Sheet1"
            wb.save(path)
            wb.close()
            info = get_data_validation_info(path, "Sheet1")
            self.assertEqual(
                info, {"sheet_name": "Sheet1", "validation_rules": []}
            )


class Phase4RemoteTests(unittest.TestCase):
    def setUp(self):
        self._old_env = {
            "EXCEL_MCP_ALLOWED_URL_HOSTS": os.environ.get(
                "EXCEL_MCP_ALLOWED_URL_HOSTS"
            ),
            "EXCEL_MCP_MAX_FILE_BYTES": os.environ.get(
                "EXCEL_MCP_MAX_FILE_BYTES"
            ),
            "EXCEL_MCP_REQUEST_TIMEOUT_SECONDS": os.environ.get(
                "EXCEL_MCP_REQUEST_TIMEOUT_SECONDS"
            ),
        }
        os.environ["EXCEL_MCP_ALLOWED_URL_HOSTS"] = ALLOWED_HOST
        os.environ["EXCEL_MCP_MAX_FILE_BYTES"] = str(1024 * 1024)
        os.environ["EXCEL_MCP_REQUEST_TIMEOUT_SECONDS"] = "30"

    def tearDown(self):
        for key, value in self._old_env.items():
            if value is None:
                os.environ.pop(key, None)
            else:
                os.environ[key] = value

    def _sample_bytes(self) -> bytes:
        stream = io.BytesIO()
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        data = [
            ["Name", "Score", "Region"],
            ["Ada", 10, "West"],
            ["Bob", 20, "East"],
            ["Cara", 30, "West"],
        ]
        for r, row in enumerate(data, start=1):
            for c, value in enumerate(row, start=1):
                ws.cell(r, c, value)
        wb.save(stream)
        wb.close()
        return stream.getvalue()

    def _run_mutation(self, **kwargs):
        source = self._sample_bytes()
        put_response = MagicMock(status_code=200, is_redirect=False)
        fake_client = _FakeClient(
            get_response=_FakeStreamResponse(source),
            put_response=put_response,
        )
        with patch.object(remote, "_create_http_client", return_value=fake_client):
            result = remote.execute_workbook_job(
                input_download_url=_url("/in.xlsx"),
                output_upload_url=_url("/out.xlsx"),
                sheet_name="Sheet1",
                **kwargs,
            )
        return result, fake_client

    def test_insert_rows_shifts_down(self):
        result, fake_client = self._run_mutation(
            operation="insert_rows", start_row=2, count=2
        )
        self.assertTrue(result["success"])
        self.assertEqual(result["operation"], "insert_rows")
        self.assertTrue(result["output_uploaded"])
        self.assertIsNone(result["data"])
        self.assertEqual(len(fake_client.get_calls), 1)
        self.assertEqual(len(fake_client.put_calls), 1)
        wb = load_workbook(io.BytesIO(fake_client.put_calls[0]["body"]))
        try:
            ws = wb["Sheet1"]
            self.assertEqual(ws["A1"].value, "Name")
            self.assertIsNone(ws["A2"].value)
            self.assertIsNone(ws["A3"].value)
            self.assertEqual(ws["A4"].value, "Ada")
            self.assertEqual(ws["A5"].value, "Bob")
        finally:
            wb.close()

    def test_insert_columns_shifts_right(self):
        result, fake_client = self._run_mutation(
            operation="insert_columns", start_col=2, count=1
        )
        self.assertTrue(result["success"])
        wb = load_workbook(io.BytesIO(fake_client.put_calls[0]["body"]))
        try:
            ws = wb["Sheet1"]
            self.assertEqual(ws["A1"].value, "Name")
            self.assertIsNone(ws["B1"].value)
            self.assertEqual(ws["C1"].value, "Score")
            self.assertEqual(ws["D1"].value, "Region")
        finally:
            wb.close()

    def test_delete_sheet_rows(self):
        result, fake_client = self._run_mutation(
            operation="delete_sheet_rows", start_row=2, count=1
        )
        self.assertTrue(result["success"])
        wb = load_workbook(io.BytesIO(fake_client.put_calls[0]["body"]))
        try:
            ws = wb["Sheet1"]
            self.assertEqual(ws["A1"].value, "Name")
            self.assertEqual(ws["A2"].value, "Bob")
            self.assertEqual(ws["A3"].value, "Cara")
        finally:
            wb.close()

    def test_delete_sheet_columns(self):
        result, fake_client = self._run_mutation(
            operation="delete_sheet_columns", start_col=2, count=1
        )
        self.assertTrue(result["success"])
        wb = load_workbook(io.BytesIO(fake_client.put_calls[0]["body"]))
        try:
            ws = wb["Sheet1"]
            self.assertEqual(ws["A1"].value, "Name")
            self.assertEqual(ws["B1"].value, "Region")
            self.assertEqual(ws["A2"].value, "Ada")
            self.assertEqual(ws["B2"].value, "West")
        finally:
            wb.close()

    def test_count_defaults_to_one(self):
        result, fake_client = self._run_mutation(
            operation="insert_rows", start_row=2
        )
        self.assertTrue(result["success"])
        wb = load_workbook(io.BytesIO(fake_client.put_calls[0]["body"]))
        try:
            self.assertIsNone(wb["Sheet1"]["A2"].value)
            self.assertEqual(wb["Sheet1"]["A3"].value, "Ada")
        finally:
            wb.close()

    def test_rejects_bool_and_float_and_string(self):
        for bad in (True, 1.5, "2"):
            result, fake_client = self._run_mutation(
                operation="insert_rows", start_row=bad, count=1
            )
            self.assertFalse(result["success"], bad)
            self.assertIn("positive integer", result["error"])
            self.assertEqual(fake_client.put_calls, [])

    def test_row_ops_reject_start_col(self):
        result, fake_client = self._run_mutation(
            operation="insert_rows", start_row=2, start_col=1
        )
        self.assertFalse(result["success"])
        self.assertIn("start_col", result["error"])
        self.assertEqual(fake_client.get_calls, [])

    def test_column_ops_reject_start_row(self):
        result, fake_client = self._run_mutation(
            operation="insert_columns", start_col=2, start_row=1
        )
        self.assertFalse(result["success"])
        self.assertIn("start_row", result["error"])
        self.assertEqual(fake_client.get_calls, [])

    def test_delete_start_beyond_used_bounds_fails(self):
        result, fake_client = self._run_mutation(
            operation="delete_sheet_rows", start_row=100, count=1
        )
        self.assertFalse(result["success"])
        self.assertIn("exceeds worksheet bounds", result["error"])
        self.assertEqual(fake_client.put_calls, [])

    def test_start_beyond_excel_hard_limit_fails_before_download(self):
        fake_client = _FakeClient()
        with patch.object(remote, "_create_http_client", return_value=fake_client):
            result = remote.execute_workbook_job(
                operation="insert_rows",
                input_download_url=_url("/in.xlsx"),
                output_upload_url=_url("/out.xlsx"),
                sheet_name="Sheet1",
                start_row=1_048_577,
                count=1,
            )
        self.assertFalse(result["success"])
        self.assertEqual(fake_client.get_calls, [])


class Phase4LocalHelperTests(unittest.TestCase):
    def _write_sample(self, path: str):
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        for r, row in enumerate(
            [
                ["Name", "Score", "Region"],
                ["Ada", 10, "West"],
                ["Bob", 20, "East"],
            ],
            start=1,
        ):
            for c, value in enumerate(row, start=1):
                ws.cell(r, c, value)
        wb.save(path)
        wb.close()

    def test_insert_and_delete_helpers(self):
        from excel_mcp.sheet import (
            delete_cols,
            delete_rows,
            insert_cols,
            insert_row,
        )

        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "struct.xlsx")
            self._write_sample(path)
            insert_row(path, "Sheet1", 2, 1)
            insert_cols(path, "Sheet1", 2, 1)
            delete_rows(path, "Sheet1", 2, 1)
            delete_cols(path, "Sheet1", 2, 1)
            wb = load_workbook(path)
            try:
                ws = wb["Sheet1"]
                self.assertEqual(ws["A1"].value, "Name")
                self.assertEqual(ws["B1"].value, "Score")
                self.assertEqual(ws["A2"].value, "Ada")
            finally:
                wb.close()

    def test_delete_count_past_used_area(self):
        from excel_mcp.sheet import delete_rows

        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "del.xlsx")
            self._write_sample(path)
            # start at last used row with count beyond used area
            delete_rows(path, "Sheet1", 3, 5)
            wb = load_workbook(path)
            try:
                self.assertEqual(wb["Sheet1"]["A1"].value, "Name")
                self.assertEqual(wb["Sheet1"]["A2"].value, "Ada")
                self.assertIsNone(wb["Sheet1"]["A3"].value)
            finally:
                wb.close()


class ExistingToolsStillWork(unittest.TestCase):
    def test_create_workbook_local_path(self):
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "local.xlsx")
            result = create_workbook(path)
            try:
                self.assertTrue(os.path.exists(path))
                self.assertIn("Sheet1", result["workbook"].sheetnames)
            finally:
                result["workbook"].close()


if __name__ == "__main__":
    unittest.main()
